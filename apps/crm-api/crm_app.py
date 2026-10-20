"""TFSC Sales CRM. Run uvicorn crm_app:app --host 127.0.0.1 --port 8766."""
from __future__ import annotations
import os, re, json, time, secrets, hashlib, base64, csv, io, hmac
from datetime import datetime, date, timedelta, timezone
from zoneinfo import ZoneInfo
from contextlib import asynccontextmanager
from urllib.parse import urlsplit
from pathlib import Path
import phonenumbers
import httpx
from fastapi import FastAPI, Request, HTTPException, UploadFile, File
from fastapi.responses import JSONResponse, FileResponse, Response
from sqlalchemy import select, or_, and_, func, case, delete
from sqlalchemy.exc import IntegrityError
import crm_db as db
import crm_security as sec
import crm_reporting as reporting
import crm_followups as followups

VERSION='1.0.1'
ORIGIN=os.getenv('CRM_ORIGIN','http://127.0.0.1:8766').rstrip('/')
PACIFIC=ZoneInfo('America/Los_Angeles')
if not db.LOCAL and (not ORIGIN.startswith('https://') or db.URL.startswith('sqlite')):raise RuntimeError('Hosted CRM requires HTTPS CRM_ORIGIN and PostgreSQL CRM_DATABASE_URL.')
@asynccontextmanager
async def lifespan(app):
 db.init_db()
 with db.engine.begin() as c:reporting.initialize(c)
 yield
app=FastAPI(title='TFSC Sales CRM',version=VERSION,lifespan=lifespan,docs_url='/api/docs' if db.LOCAL else None,redoc_url=None)

@app.middleware('http')
async def boundary(request,call_next):
 allowed={urlsplit(ORIGIN).netloc}
 if db.LOCAL:allowed.add('localhost:'+str(urlsplit(ORIGIN).port or 80))
 if request.headers.get('host','') not in allowed:return JSONResponse({'error':'Invalid application host.'},403)
 if request.method not in ('GET','HEAD','OPTIONS'):
  if request.headers.get('origin') not in (None,ORIGIN) or request.headers.get('sec-fetch-site')=='cross-site':return JSONResponse({'error':'Request origin was rejected.'},403)
  size=0;chunks=[]
  async for chunk in request.stream():
   size+=len(chunk)
   if size>32_000_000:return JSONResponse({'error':'Request exceeds the 32 MB limit.'},413)
   chunks.append(chunk)
  request._body=b''.join(chunks)
 try:r=await call_next(request)
 except Exception:
  import logging;logging.exception('CRM request failed');r=JSONResponse({'error':'The request could not be completed. Please retry.'},500)
 r.headers.update({'Cache-Control':'no-store','X-Content-Type-Options':'nosniff','X-Frame-Options':'DENY','Referrer-Policy':'no-referrer','Content-Security-Policy':"default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; object-src 'none'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'"})
 if not db.LOCAL:r.headers['Strict-Transport-Security']='max-age=31536000'
 return r
@app.exception_handler(HTTPException)
async def invalid_request(request,e):return JSONResponse({'error':e.detail},e.status_code)
@app.exception_handler(ValueError)
async def invalid_value(request,e):return JSONResponse({'error':str(e)},400)
@app.exception_handler(IntegrityError)
async def conflict(request,e):return JSONResponse({'error':'This change conflicts with an existing record. Refresh and review it.'},409)

def clean(v,label,required=False,limit=500):
 if v is None:v=''
 if not isinstance(v,str):raise ValueError(label+' must be text.')
 v=v.strip()
 if (required and not v) or len(v)>limit:raise ValueError(label+' is missing or too long.')
 if any(ord(x)<32 and x not in '\n\t\r' for x in v):raise ValueError(label+' contains control characters.')
 return v
def email(v,required=False):
 v=clean(v,'Email',required,250)
 if v and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',v):raise ValueError('Enter a valid email address without internal spaces.')
 if v:
  local,domain=v.rsplit('@',1);v=local+'@'+domain.lower()
 return v
def today():return datetime.now(PACIFIC).date().isoformat()
def iso_date(v):
 v=v or today()
 try:
  if isinstance(v,(int,float)) or (isinstance(v,str) and re.fullmatch(r'\d{5}(?:\.0+)?',v)):
   n=float(v)
   if n<1 or n>100000 or not n.is_integer():raise ValueError()
   return (date(1899,12,30)+timedelta(days=int(n))).isoformat()
  if isinstance(v,str) and '/' in v:return datetime.strptime(v,'%m/%d/%Y').date().isoformat()
  return date.fromisoformat(str(v)).isoformat()
 except (ValueError,OverflowError):raise ValueError('Use YYYY-MM-DD or an unambiguous U.S. MM/DD/YYYY date.')
def tri(v):
 if v is None or v=='':return None
 if type(v) is bool:return v
 if isinstance(v,str) and v.casefold() in ('yes','no','unset'):return {'yes':True,'no':False,'unset':None}[v.casefold()]
 raise ValueError('Outcomes must be Yes, No or Unset.')
def normal_phone(v):
 raw=clean(v,'Phone',False,100)
 if not raw:return '',None,None,[]
 try:
  p=phonenumbers.parse(raw,'US')
  if p.country_code!=1 or not phonenumbers.is_possible_number(p):raise ValueError()
  return phonenumbers.format_number(p,phonenumbers.PhoneNumberFormat.NATIONAL),phonenumbers.format_number(p,phonenumbers.PhoneNumberFormat.E164),p.extension or None,[]
 except Exception:return raw,None,None,['Phone needs review; it could not be normalized as a North American number.']
def normalized(payload,source='manual'):
 name=clean(payload.get('provider_name',payload.get('name','')),'Provider or practice name',True,200)
 phone,key,extension,warnings=normal_phone(payload.get('phone',''))
 em=email(payload.get('email',''))
 outcome={k:tri(payload.get(k)) for k in ('called','connected','text_sent','email_sent')}
 if outcome['connected'] is True and outcome['called'] is not True:raise ValueError('Connected requires Called = Yes. Keep inconsistent imported rows for review.')
 nextdate=clean(payload.get('next_followup',''),'Next follow-up',False,40) or None
 if nextdate:
  try:
   if len(nextdate)==10:date.fromisoformat(nextdate)
   else:
    d=datetime.fromisoformat(nextdate.replace('Z','+00:00'))
    if d.tzinfo is None:raise ValueError()
  except ValueError:raise ValueError('Follow-up must be a date or an ISO time with an explicit offset.')
 return dict(provider_name=name,phone=phone,phone_key=key,extension=extension,email=em,address=clean(payload.get('address',''),'Address',False,1000),business_date=iso_date(payload.get('business_date')),next_followup=nextdate,remarks=clean(payload.get('remarks',''),'Remarks',False,6000),**outcome),warnings
def body_hash(payload):return hashlib.sha256(json.dumps(payload,sort_keys=True,separators=(',',':'),ensure_ascii=False).encode()).hexdigest()
def get_record(c,a,table,id,owner='owner_id'):
 r=c.execute(select(table).where(table.c.id==id,table.c.org_id==a['org_id'],sec.scope(c,a,table.c[owner]))).mappings().first()
 if not r:raise HTTPException(404,'Record not found.')
 return dict(r)
def version(payload,r):
 if payload.get('version')!=r['version']:raise HTTPException(409,'This record has changed. Refresh before editing.')
def page(query,table,request,c):
 limit=min(max(int(request.query_params.get('limit','50')),1),200)
 cursor=request.query_params.get('cursor')
 if cursor:
  try:
   stamp,id=json.loads(base64.urlsafe_b64decode(cursor.encode()).decode());query=query.where(or_(table.c.created<float(stamp),and_(table.c.created==float(stamp),table.c.id<id)))
  except Exception:raise ValueError('Invalid page cursor.')
 rows=[dict(r) for r in c.execute(query.order_by(table.c.created.desc(),table.c.id.desc()).limit(limit+1)).mappings()]
 next_cursor=None
 if len(rows)>limit:
  rows=rows[:limit];r=rows[-1];next_cursor=base64.urlsafe_b64encode(json.dumps([r['created'],r['id']]).encode()).decode()
 return rows,next_cursor
def stage_quality(r,stage=None):
 if r['called'] is not True:return ('Research' if r['first_committed_at'] is None else 'No call',None)
 if stage=='signed_up':return 'Sign Up',100
 if stage in ('strong_follow_up_confirmed','qualified_lead','meeting'):return 'Strong Follow-Up',50
 if r['connected'] and (r['text_sent'] or r['email_sent']):return 'Touch + Connect',30
 if r['connected']:return 'Connected',20
 return 'Attempted',10
def activity_dto(r,a,stage=None):
 keys=['id','owner_id','provider_name','phone','email','address','business_date','called','connected','text_sent','email_sent','next_followup','remarks','created','updated','first_committed_at','edit_deadline','version','source']
 out={k:r[k] for k in keys};out['stage'],out['quality']=stage_quality(r,stage)
 out['locked']=bool(r['edit_deadline'] and time.time()>=r['edit_deadline']);out['can_correct']=sec.rank(a)>=1
 if sec.can_routing(a):out.update(domain_id=r['domain_id'],project_id=r['project_id'])
 return out
def session_response(c,u):
 token=secrets.token_urlsafe(32);csrf=secrets.token_urlsafe(32)
 c.execute(delete(db.sessions).where(db.sessions.c.expires<time.time()))
 c.execute(db.sessions.insert().values(token=sec.digest(token),user_id=u['id'],csrf=csrf,expires=time.time()+28800))
 response=JSONResponse({'user':sec.public(u),'csrf':csrf})
 response.set_cookie('crm_session',token,httponly=True,secure=not db.LOCAL,samesite='lax',max_age=28800,path='/')
 return response

@app.get('/api/health')
def health():
 with db.engine.connect() as c:c.execute(select(1))
 return {'app':'TFSC Sales CRM','version':VERSION,'mode':'local' if db.LOCAL else 'hosted'}
@app.get('/api/auth/status')
def auth_status():
 with db.engine.connect() as c:return {'setup_required':c.execute(select(db.users.c.id).limit(1)).first() is None}
@app.post('/api/auth/setup')
def setup(payload:dict,request:Request):
 with db.engine.begin() as c:
  if not sec.rate_limit(c,'setup:'+request.client.host,5):raise HTTPException(429,'Too many setup attempts.')
 key=db.DATA/'setup-key.txt'
 if not key.exists() or not secrets.compare_digest(str(payload.get('setup_key','')),key.read_text().strip()):raise HTTPException(403,'Enter the setup key from the CRM data folder.')
 name=clean(payload.get('name'),'Name',True,150);em=email(payload.get('email'),True).casefold();pw=sec.password_hash(payload.get('password'))
 with db.engine.begin() as c:
  if c.execute(select(db.users.c.id).limit(1)).first():raise HTTPException(409,'Workspace is already configured.')
  c.execute(db.settings.insert().values(key='setup_complete',value=True))
  u=dict(id=db.uid(),org_id=db.uid(),name=name,email=em,password=pw,role='super_admin',manager_id=None,agent_id='ADMIN',active=True,routing_view=True,domain_id='scribe_align',project_id=None,agreements_user_id=None,created=time.time(),version=1,photo=None,reset_required=False)
  c.execute(db.users.insert().values(**u));db.log(c,u,'workspace.initialized');r=session_response(c,u)
 key.unlink(missing_ok=True);return r
@app.post('/api/auth/login')
def login(payload:dict,request:Request):
 em=str(payload.get('email','')).strip().casefold()[:200]
 with db.engine.begin() as c:
  allowed=sec.rate_limit(c,'ip:'+request.client.host,60) and sec.rate_limit(c,'login:'+em,10)
 if not allowed:raise HTTPException(429,'Too many sign-in attempts. Try again in 15 minutes.')
 with db.engine.begin() as c:
  u=c.execute(select(db.users).where(db.users.c.email==em,db.users.c.active==True)).mappings().first()
  if not u or not sec.verify(u['password'],str(payload.get('password',''))[:200]):raise HTTPException(401,'Email or password is incorrect.')
  db.log(c,u,'user.signed_in');return session_response(c,u)
@app.get('/api/auth/me')
def me(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c,True);return {'user':sec.public(a),'csrf':a['_csrf']}
@app.post('/api/auth/logout')
def logout(request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c,True);c.execute(delete(db.sessions).where(db.sessions.c.token==sec.digest(request.cookies['crm_session'])));db.log(c,a,'user.signed_out')
 r=JSONResponse({'ok':True});r.delete_cookie('crm_session');return r
@app.post('/api/auth/password')
def password(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c,True)
  if not sec.verify(a['password'],str(payload.get('current_password',''))):raise HTTPException(403,'Current password is incorrect.')
  c.execute(db.users.update().where(db.users.c.id==a['id']).values(password=sec.password_hash(payload.get('password')),reset_required=False))
  c.execute(delete(db.sessions).where(db.sessions.c.user_id==a['id'],db.sessions.c.token!=sec.digest(request.cookies['crm_session'])));db.log(c,a,'user.password_changed')
 return {'ok':True}

@app.get('/api/config')
def config(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);p=db.policy(c)
  result={'user':sec.public(a),'capabilities':sec.capabilities(a),'settings':p,'business_date':today(),'timezone':'America/Los_Angeles','version':VERSION,'agreements_url':os.getenv('CRM_TFSC_URL','http://127.0.0.1:8765') if a['agreements_user_id'] else None}
  if sec.can_routing(a):result.update(domains=[dict(r) for r in c.execute(select(db.domains)).mappings()],projects=[dict(r) for r in c.execute(select(db.projects)).mappings()])
  return result

@app.get('/api/users')
def list_users(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);sec.require(a)
  q=select(db.users).where(db.users.c.org_id==a['org_id'],sec.scope(c,a,db.users.c.id))
  if request.query_params.get('q'):q=q.where(db.users.c.name.ilike('%'+request.query_params['q'][:100]+'%'))
  rows,cursor=page(q,db.users,request,c)
  return {'items':[sec.public(r,sec.can_routing(a)) for r in rows],'next_cursor':cursor}

def user_values(c,a,p,old=None):
 sec.require(a,2);role=p.get('role',old['role'] if old else 'sales_agent')
 if role not in sec.ROLES or (a['role']!='super_admin' and sec.ROLES[role]>=sec.rank(a)):raise HTTPException(403,'You cannot assign this role.')
 manager=p.get('manager_id',old['manager_id'] if old else a['id']) or a['id']
 parent=c.execute(select(db.users).where(db.users.c.id==manager,db.users.c.org_id==a['org_id'],db.users.c.active==True)).mappings().first()
 if not parent or not sec.allowed_user(c,a,manager) or sec.rank(parent)<=sec.ROLES[role]:raise ValueError('Choose an eligible manager within your scope.')
 if old:
  if old['id']==a['id']:raise HTTPException(403,'Use a separate administrator to change your own account permissions.')
  if manager==old['id'] or c.execute(select(db.users.c.id).where(db.users.c.id==manager,sec.scope(c,old,db.users.c.id))).first():raise ValueError('Reporting relationships cannot contain a cycle.')
 dom=p.get('domain_id',old['domain_id'] if old else 'scribe_align') or 'scribe_align';proj=p.get('project_id',old['project_id'] if old else None) or None
 if not c.execute(select(db.domains).where(db.domains.c.id==dom,db.domains.c.active==True)).first():raise ValueError('Choose an active sales domain.')
 if proj and not c.execute(select(db.projects).where(db.projects.c.id==proj,db.projects.c.active==True)).first():raise ValueError('Choose an active project.')
 routing=bool(p.get('routing_view',old['routing_view'] if old else False)) and role!='sales_agent'
 linked=p.get('agreements_user_id',old['agreements_user_id'] if old else None) or None
 if linked!=(old['agreements_user_id'] if old else None) and a['role']!='super_admin':raise HTTPException(403,'Only Super Admin can link agreements access.')
 if linked and not re.fullmatch(r'[a-f0-9]{32}',linked):raise ValueError('Use the existing TFSC user ID for an explicit account link.')
 vals=dict(name=clean(p.get('name',old['name'] if old else ''),'Name',True,150),email=email(p.get('email',old['email'] if old else ''),True).casefold(),role=role,manager_id=manager,domain_id=dom,project_id=proj,routing_view=routing,agreements_user_id=linked,agent_id=clean(p.get('agent_id',old['agent_id'] if old else ''),'Agent ID',False,80),active=bool(p.get('active',old['active'] if old else True)))
 if not old or p.get('password'):vals.update(password=sec.password_hash(p.get('password')),reset_required=bool(p.get('reset_required',True)))
 return vals
@app.post('/api/users')
def create_user(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);v=user_values(c,a,payload);v.update(id=db.uid(),org_id=a['org_id'],created=time.time(),version=1,photo=None)
  c.execute(db.users.insert().values(**v));db.log(c,a,'user.created',v['id'],{'role':v['role']},v['id']);return sec.public(v,sec.can_routing(a))
@app.patch('/api/users/{id}')
def update_user(id:str,payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);old=get_record(c,a,db.users,id,'id');version(payload,old);v=user_values(c,a,payload,old);v['version']=old['version']+1
  changed=c.execute(db.users.update().where(db.users.c.id==id,db.users.c.version==old['version']).values(**v)).rowcount
  if not changed:raise HTTPException(409,'User changed. Refresh.')
  c.execute(delete(db.sessions).where(db.sessions.c.user_id==id));db.log(c,a,'user.updated',id,{'before':{k:old[k] for k in ('role','manager_id','domain_id','project_id')},'after':{k:v[k] for k in ('role','manager_id','domain_id','project_id')}},id)
  return sec.public({**old,**v},sec.can_routing(a))

def duplicate_check(c,a,r):
 if not r.get('phone_key'):return None
 others=c.execute(select(db.activities.c.id,db.activities.c.owner_id,db.activities.c.next_followup,db.leads.c.stage).outerjoin(db.leads,db.leads.c.activity_id==db.activities.c.id).where(db.activities.c.org_id==a['org_id'],db.activities.c.phone_key==r['phone_key'],db.activities.c.id!=r['id'],db.activities.c.deleted==False).limit(20)).mappings().all()
 claim=c.execute(select(db.claims.c.owner_id).where(db.claims.c.claim_key==a['org_id']+':'+r['phone_key'])).scalar_one_or_none()
 severity='owned_lead' if claim and claim!=r['owner_id'] else None
 for o in others:
  if o['owner_id']==r['owner_id']:continue
  current='owned_lead' if o['stage'] in ('qualified_lead','meeting','signed_up') else 'strong_follow_up' if o['stage']=='strong_follow_up_confirmed' else 'follow_up' if o['next_followup'] else 'duplicate'
  ranks={'duplicate':1,'follow_up':2,'strong_follow_up':3,'owned_lead':4}
  if not severity or ranks[current]>ranks[severity]:severity=current
  exists=c.execute(select(db.conflicts.c.id).where(db.conflicts.c.activity_id==r['id'],db.conflicts.c.other_id==o['id'],db.conflicts.c.status=='open')).first()
  if not exists:c.execute(db.conflicts.insert().values(id=db.uid(),org_id=a['org_id'],activity_id=r['id'],other_id=o['id'],severity=current,status='open',created=time.time()))
 return severity
def create_activity(c,a,p,source='manual',assigned=None):
 v,warnings=normalized(p,source);now=time.time();committed=any(v[k] is not None for k in ('called','connected','text_sent','email_sent'))
 duration=60 if source=='manual' else db.policy(c)['import_edit_minutes']
 original=p.get('_original',p)
 if not isinstance(original,dict):raise ValueError('Invalid source row.')
 owner=assigned or a
 v.update(id=db.uid(),org_id=a['org_id'],owner_id=owner['id'],raw={k:original.get(k) for k in ('provider_name','phone','email','business_date','strong_followup_requested','sale_requested','historical_followup_count','last_followup','call_again')},source=source,domain_id=owner['domain_id'],project_id=owner['project_id'],created=now,updated=now,first_committed_at=now if committed else None,edit_deadline=now+duration*60 if committed else None,version=1,deleted=False)
 c.execute(db.activities.insert().values(**v));severity=duplicate_check(c,a,v)
 if v['next_followup']:c.execute(db.leads.insert().values(id=db.uid(),org_id=a['org_id'],activity_id=v['id'],owner_id=owner['id'],stage='follow_up',version=1,created=now,updated=now))
 db.log(c,a,'activity.created',v['id'],{'source':source,'committed':committed},owner['id'])
 result=activity_dto(v,a);result.update(warnings=warnings,duplicate_severity=severity);return result
@app.post('/api/activities')
def create_activity_route(payload:dict,request:Request):
 with db.engine.begin() as c:return create_activity(c,sec.actor(request,c),payload)
@app.get('/api/activities')
def list_activities(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);q=select(db.activities,db.leads.c.stage.label('lead_stage')).outerjoin(db.leads,db.leads.c.activity_id==db.activities.c.id).where(db.activities.c.org_id==a['org_id'],db.activities.c.deleted==False,sec.scope(c,a,db.activities.c.owner_id))
  if request.query_params.get('q'):
   term='%'+request.query_params['q'][:100]+'%';q=q.where(or_(db.activities.c.provider_name.ilike(term),db.activities.c.phone.ilike(term)))
  if request.query_params.get('from'):q=q.where(db.activities.c.business_date>=iso_date(request.query_params['from']))
  if request.query_params.get('to'):q=q.where(db.activities.c.business_date<=iso_date(request.query_params['to']))
  rows,cursor=page(q,db.activities,request,c)
  return {'items':[activity_dto(r,a,r.get('lead_stage')) for r in rows],'next_cursor':cursor}
@app.patch('/api/activities/{id}')
def edit_activity(id:str,payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);old=get_record(c,a,db.activities,id);version(payload,old)
  if old['deleted']:raise HTTPException(409,'This record is archived.')
  locked=bool(old['edit_deadline'] and time.time()>=old['edit_deadline'])
  if locked and sec.rank(a)<1:raise HTTPException(403,'The edit window has closed. Ask your supervisor for a correction.')
  reason=clean(payload.get('reason',''),'Correction reason',locked or old['owner_id']!=a['id'],1000)
  values,warnings=normalized({**old,**payload},old['source']);now=time.time()
  owned=c.execute(select(db.claims.c.claim_key).join(db.leads,db.leads.c.id==db.claims.c.lead_id).where(db.leads.c.activity_id==id)).first()
  if owned and any(values[k]!=old[k] for k in ('provider_name','phone_key')):raise HTTPException(409,'This activity supports an owned lead. Resolve identity and ownership before changing its name or phone.')
  if not old['first_committed_at'] and any(values[k] is not None for k in ('called','connected','text_sent','email_sent')):
   values['first_committed_at']=now;values['edit_deadline']=now+(60 if old['source']=='manual' else db.policy(c)['import_edit_minutes'])*60
  values.update(version=old['version']+1,updated=now)
  # CAS prevents concurrent overwrites. Deadline is re-evaluated in the UPDATE itself.
  q=db.activities.update().where(db.activities.c.id==id,db.activities.c.version==old['version'])
  if sec.rank(a)<1:q=q.where(or_(db.activities.c.edit_deadline==None,db.activities.c.edit_deadline>time.time()))
  if c.execute(q.values(**values)).rowcount!=1:raise HTTPException(409,'The record changed or its edit window closed.')
  result={**old,**values};severity=duplicate_check(c,a,result)
  db.log(c,a,'activity.corrected' if locked else 'activity.updated',id,{'reason':reason,'before':{k:old[k] for k in values if k not in ('version','updated')},'after':values},old['owner_id'])
  lead=c.execute(select(db.leads).where(db.leads.c.activity_id==id)).mappings().first()
  if not lead and result['next_followup']:c.execute(db.leads.insert().values(id=db.uid(),org_id=a['org_id'],activity_id=id,owner_id=old['owner_id'],stage='follow_up',version=1,created=now,updated=now))
  out=activity_dto(result,a,lead['stage'] if lead else None);out.update(warnings=warnings,duplicate_severity=severity);return out
@app.delete('/api/activities/{id}')
def delete_activity(id:str,payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a);r=get_record(c,a,db.activities,id);version(payload,r);reason=clean(payload.get('reason'),'Reason',True,1000)
  if c.execute(db.activities.update().where(db.activities.c.id==id,db.activities.c.version==r['version']).values(deleted=True,version=r['version']+1,updated=time.time())).rowcount!=1:raise HTTPException(409,'Record changed.')
  db.log(c,a,'activity.archived',id,{'reason':reason},r['owner_id']);return {'ok':True}

@app.post('/api/refine')
def refine(payload:dict,request:Request):
 with db.engine.connect() as c:sec.actor(request,c)
 rows=payload.get('rows',[])
 if not isinstance(rows,list) or not 1<=len(rows)<=1000:raise ValueError('Refine 1 to 1,000 rows per batch.')
 result=[]
 for i,r in enumerate(rows):
  try:
   if not isinstance(r,dict):raise ValueError('Each row must be an object.')
   v,warnings=normalized(r,'import');v.pop('phone_key',None);v.pop('extension',None);v['_original']={k:r.get(k) for k in ('provider_name','phone','email','business_date','strong_followup_requested','sale_requested','historical_followup_count','last_followup','call_again')}
   if r.get('strong_followup_requested') or r.get('sale_requested'):warnings.append('Historical qualification/sign-up claim retained as source evidence; a supervisor must confirm it in the lead workflow.')
   result.append({'row':i+1,'original':r,'refined':v,'warnings':warnings,'valid':True})
  except (ValueError,HTTPException) as e:result.append({'row':i+1,'original':r,'error':str(e),'valid':False})
 return {'items':result,'valid':sum(r['valid'] for r in result),'invalid':sum(not r['valid'] for r in result)}
@app.post('/api/import')
def import_rows(payload:dict,request:Request):
 rows=payload.get('rows',[]);key=clean(payload.get('idempotency_key'),'Import key',True,80)
 if not isinstance(rows,list) or not 1<=len(rows)<=1000:raise ValueError('Import 1 to 1,000 rows per batch.')
 body=body_hash({'rows':rows,'owner_id':payload.get('owner_id')})
 with db.engine.begin() as c:
  a=sec.actor(request,c);existing=c.execute(select(db.imports).where(db.imports.c.owner_id==a['id'],db.imports.c.request_key==key)).mappings().first()
  if existing:
   if existing['body_hash']!=body:raise HTTPException(409,'This import key belongs to different data.')
   result=dict(existing['result'])
   if not sec.can_routing(a):result['items']=[{k:v for k,v in r.items() if k not in ('domain_id','project_id')} for r in result['items']]
   return result
  assigned=None
  if payload.get('owner_id') and payload['owner_id']!=a['id']:
   sec.require(a)
   if not sec.allowed_user(c,a,payload['owner_id']):raise HTTPException(403,'Import owner is outside your scope.')
   assigned=c.execute(select(db.users).where(db.users.c.id==payload['owner_id'],db.users.c.active==True)).mappings().first()
   if not assigned:raise ValueError('Choose an active import owner.')
  for row in rows:normalized(row,'import')
  items=[create_activity(c,a,r,'import',assigned) for r in rows]
  result={'id':db.uid(),'items':items,'saved':len(items),'skipped':0,'errors':[]}
  c.execute(db.imports.insert().values(id=result['id'],org_id=a['org_id'],owner_id=a['id'],request_key=key,body_hash=body,result=result,created=time.time()));db.log(c,a,'import.completed',result['id'],{'count':len(items)});return result

@app.post('/api/import/file-preview')
async def preview_workbook(request:Request,file:UploadFile=File(...)):
 with db.engine.connect() as c:sec.actor(request,c)
 content=await file.read()
 if len(content)>30_000_000:raise HTTPException(413,'Use a workbook smaller than 30 MB.')
 import zipfile
 from openpyxl import load_workbook
 try:
  with zipfile.ZipFile(io.BytesIO(content)) as z:
   if sum(f.file_size for f in z.infolist())>550_000_000 or len(z.infolist())>3000:raise ValueError('Workbook is too large after decompression.')
   if any('vbaProject' in f.filename for f in z.infolist()):raise ValueError('Macro-enabled workbooks are not accepted.')
  book=load_workbook(io.BytesIO(content),read_only=True,data_only=True,keep_links=False)
  sheet_name=request.query_params.get('sheet');sheets=[book[sheet_name]] if sheet_name and sheet_name in book.sheetnames else [s for s in book if s.title not in ('Dashboard','Validations') and s.sheet_state=='visible']
  accepted=[];source_sheet=None
  for sheet in sheets:
   for row in sheet.iter_rows(values_only=True):
    if any(v is not None for v in row[:4]):
     if source_sheet is None:source_sheet=sheet.title
     accepted.append([v.isoformat()[:10] if isinstance(v,(datetime,date)) else '' if v is None else str(v) for v in row[:17]])
     if len(accepted)>1001:break
   if accepted:break
  book.close();stream=io.StringIO();writer=csv.writer(stream,delimiter='\t',lineterminator='\n');writer.writerows(accepted[:1001]);return {'tsv':stream.getvalue(),'source_sheet':source_sheet,'has_more':len(accepted)>1001,'method':'Preview of first populated visible data sheet; no rows saved.'}
 except (zipfile.BadZipFile,KeyError):raise ValueError('Use a valid XLSX workbook.')

@app.post('/api/profile/photo')
def upload_photo(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c)
  try:
   from PIL import Image
   raw=base64.b64decode(str(payload.get('image','')).split(',')[-1],validate=True)
   if len(raw)>2_000_000:raise ValueError()
   im=Image.open(io.BytesIO(raw))
   if im.width*im.height>16_000_000:raise ValueError()
   im.load();im=im.convert('RGB');im.thumbnail((256,256));out=io.BytesIO();im.save(out,format='JPEG',quality=88)
  except Exception:raise ValueError('Use a PNG or JPEG smaller than 2 MB and 16 megapixels.')
  c.execute(db.users.update().where(db.users.c.id==a['id']).values(photo=base64.b64encode(out.getvalue()).decode()));db.log(c,a,'profile.photo_updated');return {'ok':True}
@app.get('/api/profile/photo/{id}')
def photo(id:str,request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c)
  if not sec.allowed_user(c,a,id):raise HTTPException(404,'Photo not found.')
  value=c.execute(select(db.users.c.photo).where(db.users.c.id==id)).scalar_one_or_none()
  if not value:raise HTTPException(404,'Photo not found.')
  return Response(base64.b64decode(value),media_type='image/jpeg')

@app.get('/api/followups')
def delegated_followups(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c)
  delegated=select(db.members.c.user_id).join(db.teams,db.teams.c.id==db.members.c.team_id).where(db.teams.c.org_id==a['org_id'],db.teams.c.leader_id==a['id'])
  visible=or_(sec.scope(c,a,db.leads.c.owner_id),db.leads.c.owner_id.in_(delegated))
  q=select(db.leads,db.activities.c.provider_name,db.activities.c.phone,db.activities.c.email,func.coalesce(followups.tasks.c.due_date,db.activities.c.next_followup).label('next_followup')).join(db.activities,db.leads.c.activity_id==db.activities.c.id).outerjoin(followups.tasks,followups.tasks.c.lead_id==db.leads.c.id).where(db.leads.c.org_id==a['org_id'],visible,db.activities.c.deleted==False,db.leads.c.stage.in_(['follow_up','strong_follow_up_pending','strong_follow_up_confirmed']))
  if request.query_params.get('q'):q=q.where(db.activities.c.provider_name.ilike('%'+request.query_params['q'][:100]+'%'))
  rows,cursor=page(q,db.leads,request,c)
  for r in rows:r.pop('org_id',None)
  return {'items':rows,'next_cursor':cursor}

@app.get('/api/leads')
@app.get('/api/clients')
def list_leads(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);q=select(db.leads,db.activities.c.provider_name,db.activities.c.phone,db.activities.c.email,db.activities.c.address,db.activities.c.next_followup).join(db.activities,db.leads.c.activity_id==db.activities.c.id).where(db.leads.c.org_id==a['org_id'],db.activities.c.deleted==False,sec.scope(c,a,db.leads.c.owner_id))
  if request.url.path.endswith('/clients'):q=q.where(db.leads.c.stage=='signed_up')
  if request.query_params.get('stage'):q=q.where(db.leads.c.stage==request.query_params['stage'])
  if request.query_params.get('q'):q=q.where(db.activities.c.provider_name.ilike('%'+request.query_params['q'][:100]+'%'))
  rows,cursor=page(q,db.leads,request,c)
  return {'items':rows,'next_cursor':cursor}
@app.post('/api/activities/{id}/followup')
def create_followup(id:str,payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);r=get_record(c,a,db.activities,id)
  existing=c.execute(select(db.leads).where(db.leads.c.activity_id==id)).mappings().first()
  if existing:return dict(existing)
  v=dict(id=db.uid(),org_id=a['org_id'],activity_id=id,owner_id=r['owner_id'],stage='follow_up',version=1,created=time.time(),updated=time.time())
  c.execute(db.leads.insert().values(**v));db.log(c,a,'followup.created',v['id'],owner_id=r['owner_id']);return v
@app.post('/api/leads/{id}/transition')
def lead_transition(id:str,payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);r=get_record(c,a,db.leads,id);version(payload,r);action=payload.get('action');reason=clean(payload.get('reason',''),'Reason',action!='request_strong',2000)
  mapping={'request_strong':('strong_follow_up_pending',{'follow_up'}),'confirm_strong':('strong_follow_up_confirmed',{'follow_up','strong_follow_up_pending'}),'qualify':('qualified_lead',{'strong_follow_up_confirmed'}),'confirm_sale':('signed_up',{'qualified_lead','meeting'}),'lost':('lost',{'follow_up','strong_follow_up_pending','strong_follow_up_confirmed','qualified_lead','meeting'}),'reopen':('follow_up',{'lost'}),'reverse_sale':('qualified_lead',{'signed_up'})}
  if action not in mapping:raise ValueError('Unknown transition.')
  target,allowed=mapping[action]
  if r['stage'] not in allowed:raise HTTPException(409,'This transition is not allowed from the current stage.')
  if action!='request_strong':sec.require(a)
  owner=payload.get('assigned_user_id') or r['owner_id']
  if owner!=r['owner_id'] and action!='qualify':raise ValueError('Owner can be assigned only during qualification; use the audited reassignment route later.')
  if not sec.allowed_user(c,a,owner):raise HTTPException(403,'Assigned user is outside your scope.')
  active=c.execute(select(db.users.c.id).where(db.users.c.id==owner,db.users.c.active==True)).first()
  if not active:raise ValueError('Assigned user is inactive.')
  v={'stage':target,'version':r['version']+1,'updated':time.time(),'owner_id':owner}
  if action=='qualify':
   activity=c.execute(select(db.activities).where(db.activities.c.id==r['activity_id'])).mappings().one()
   if not activity['phone_key']:raise ValueError('Normalize and review the phone before qualifying this lead.')
   c.execute(db.claims.insert().values(claim_key=a['org_id']+':'+activity['phone_key'],lead_id=id,owner_id=owner,created=time.time()));v['qualified_by']=a['id']
  if action=='confirm_sale':v['sale_confirmed_at']=time.time()
  if action=='reverse_sale':v['sale_confirmed_at']=None
  if action=='lost':c.execute(delete(db.claims).where(db.claims.c.lead_id==id))
  if c.execute(db.leads.update().where(db.leads.c.id==id,db.leads.c.version==r['version']).values(**v)).rowcount!=1:raise HTTPException(409,'Lead changed. Refresh before deciding.')
  db.log(c,a,'lead.'+action,id,{'reason':reason,'before':r['stage'],'after':target,'old_owner':r['owner_id'],'new_owner':owner},r['owner_id'])
  if action=='confirm_sale':
   winner=c.execute(select(db.users.c.name).where(db.users.c.id==owner)).scalar_one()
   c.execute(db.outbox.insert().values(id=db.uid(),org_id=a['org_id'],owner_id=owner,event_type='sale.celebrate',payload={'name':winner,'supervisor':a['name'],'message':'Congratulations on a successful sale!'},state='published',attempts=0,available_at=time.time(),created=time.time()))
  return {**r,**v}
@app.post('/api/leads/{id}/assign')
def reassign(id:str,payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a);r=get_record(c,a,db.leads,id);version(payload,r);reason=clean(payload.get('reason'),'Reason',True,1000);owner=payload.get('assigned_user_id')
  if not sec.allowed_user(c,a,owner):raise HTTPException(403,'User is outside your scope.')
  if c.execute(db.leads.update().where(db.leads.c.id==id,db.leads.c.version==r['version']).values(owner_id=owner,version=r['version']+1,updated=time.time())).rowcount!=1:raise HTTPException(409,'Lead changed.')
  c.execute(db.claims.update().where(db.claims.c.lead_id==id).values(owner_id=owner));db.log(c,a,'lead.reassigned',id,{'reason':reason,'old_owner':r['owner_id'],'new_owner':owner,'incentive_decision':clean(payload.get('incentive_decision'),'Incentive attribution decision',True,1000)},r['owner_id']);return {'ok':True}

@app.get('/api/conflicts')
def list_conflicts(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);sec.require(a)
  eligible=select(db.activities.c.id).where(sec.scope(c,a,db.activities.c.owner_id),db.activities.c.org_id==a['org_id'])
  q=select(db.conflicts).where(db.conflicts.c.org_id==a['org_id'],db.conflicts.c.status=='open',db.conflicts.c.activity_id.in_(eligible))
  rows,cursor=page(q,db.conflicts,request,c)
  for r in rows:
   if not c.execute(select(db.activities.c.id).where(db.activities.c.id==r['other_id'],sec.scope(c,a,db.activities.c.owner_id))).first():r['other_id']=None;r['requires_higher_review']=True
  return {'items':rows,'next_cursor':cursor}
@app.post('/api/conflicts/{id}/resolve')
def resolve_conflict(id:str,payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a);r=c.execute(select(db.conflicts).where(db.conflicts.c.id==id,db.conflicts.c.org_id==a['org_id'])).mappings().first()
  if not r:raise HTTPException(404,'Conflict not found.')
  get_record(c,a,db.activities,r['activity_id']);get_record(c,a,db.activities,r['other_id']);reason=clean(payload.get('reason'),'Resolution reason',True,2000)
  c.execute(db.conflicts.update().where(db.conflicts.c.id==id).values(status='resolved',reason=reason,resolved_by=a['id']));db.log(c,a,'conflict.resolved',id,{'reason':reason});return {'ok':True}

def period(request):
 end=iso_date(request.query_params.get('to'));start=iso_date(request.query_params.get('from') or end[:8]+'01')
 if start>end or (date.fromisoformat(end)-date.fromisoformat(start)).days>366:raise ValueError('Choose a reporting range of up to 366 days.')
 return start,end
def workdays(start,end):
 d=date.fromisoformat(start);end=date.fromisoformat(end);count=0
 while d<=end:
  count+=int(d.weekday()<5);d+=timedelta(days=1)
 return max(count,1)
def calc_metrics(attempts,connected,points,target,conn_target):
 q=points/(100*attempts) if attempts else 0;rate=connected/attempts if attempts else 0
 composite=.5*q+.25*min(1,attempts/max(target,1))+.25*min(1,rate/conn_target) if attempts else 0
 band='No activity' if not attempts else 'Extraordinary' if composite>=.8 else 'Satisfactory' if composite>=.65 else 'Average' if composite>=.5 else 'Below Average' if composite>=.35 else 'Poor'
 return {'attempts':attempts,'connections':connected,'connection_rate':rate,'quality':q,'composite':composite,'band':band,'target':target}
@app.get('/api/dashboard')
def dashboard(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);p=db.policy(c);start,end=period(request);ac=db.activities;le=db.leads
  score=case((le.c.stage=='signed_up',100),(le.c.stage.in_(['strong_follow_up_confirmed','qualified_lead','meeting']),50),(and_(ac.c.connected==True,or_(ac.c.text_sent==True,ac.c.email_sent==True)),30),(ac.c.connected==True,20),else_=10)
  attempts=func.sum(case((ac.c.called==True,1),else_=0));connected=func.sum(case((and_(ac.c.called==True,ac.c.connected==True),1),else_=0));points=func.sum(case((ac.c.called==True,score),else_=0))
  scope=and_(ac.c.org_id==a['org_id'],ac.c.deleted==False,sec.scope(c,a,ac.c.owner_id),ac.c.business_date>=start,ac.c.business_date<=end)
  joined=ac.outerjoin(le,le.c.activity_id==ac.c.id)
  grouped=c.execute(select(ac.c.owner_id,attempts.label('attempts'),connected.label('connections'),points.label('points'),func.min(ac.c.business_date).label('first_date')).select_from(joined).where(scope).group_by(ac.c.owner_id)).mappings().all()
  roster=c.execute(select(db.users.c.id,db.users.c.name,db.users.c.created,db.users.c.role).where(db.users.c.org_id==a['org_id'],db.users.c.active==True,sec.scope(c,a,db.users.c.id))).mappings().all();roster_map={r['id']:r for r in roster}
  groups={r['owner_id']:r for r in grouped};eligible_ids=set(groups)|{r['id'] for r in roster if r['role']=='sales_agent'}
  days=reporting.calendar(c,start,end);leaderboard=[];total_target=0;weighted_connections=0
  for owner_id in eligible_ids:
   r=groups.get(owner_id,{'attempts':0,'connections':0,'points':0,'first_date':None})
   staff=roster_map.get(owner_id,{'name':'Former team member','created':0})
   employment=datetime.fromtimestamp(staff['created'],PACIFIC).date().isoformat() if staff['created'] else start
   eligible=max(start,min(employment,r['first_date'] or employment));target,connection_target=reporting.target(days,eligible);total_target+=target;weighted_connections+=target*connection_target
   leaderboard.append({'name':staff['name'],**calc_metrics(r['attempts'],r['connections'],r['points'],target,connection_target)})
  metrics=calc_metrics(sum(r['attempts'] for r in grouped),sum(r['connections'] for r in grouped),sum(r['points'] for r in grouped),total_target,weighted_connections/total_target if total_target else .6)
  metrics['research']=c.execute(select(func.count()).select_from(ac).where(scope,ac.c.first_committed_at==None)).scalar_one()
  metrics['followups']=c.execute(select(func.count()).select_from(ac).where(scope,ac.c.next_followup!=None)).scalar_one()
  pipeline=[{'stage':r[0],'count':r[1]} for r in c.execute(select(le.c.stage,func.count()).select_from(le.join(ac,le.c.activity_id==ac.c.id)).where(scope).group_by(le.c.stage))]
  metrics['leads']=sum(r['count'] for r in pipeline if r['stage'] in ('qualified_lead','meeting'));metrics['sales']=sum(r['count'] for r in pipeline if r['stage']=='signed_up')
  trend=[dict(r) for r in c.execute(select(ac.c.business_date.label('date'),attempts.label('attempts'),connected.label('connections')).select_from(joined).where(scope).group_by(ac.c.business_date).order_by(ac.c.business_date)).mappings()]
  return {'metrics':metrics,'trend':trend,'leaderboard':sorted(leaderboard,key=lambda x:x['composite'],reverse=True)[:200],'leaderboard_total':len(leaderboard),'pipeline':pipeline,'updated_at':datetime.now(timezone.utc).isoformat(),'from':start,'to':end,'policy_version':p['version'],'target_versions':sorted(set(d[3] for d in days)),'timezone':'America/Los_Angeles','target_basis':'Active sales agents on weekdays; imported historical activity establishes earlier eligibility. No holiday or leave calendar has been configured.'}

@app.get('/api/audit')
def audit(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);q=select(db.audit).where(db.audit.c.org_id==a['org_id'],sec.scope(c,a,db.audit.c.owner_id))
  if request.query_params.get('resource_id'):q=q.where(db.audit.c.resource_id==request.query_params['resource_id'])
  rows,cursor=page(q,db.audit,request,c)
  # Never serialize private administrative before/after objects to an agent.
  for r in rows:
   r.pop('org_id',None)
   if not sec.can_routing(a):r['detail']={'reason':r['detail'].get('reason','')} if r['detail'].get('reason') else {}
  return {'items':rows,'next_cursor':cursor}
@app.get('/api/events')
def events(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);since=float(request.query_params.get('since',str(time.time()-60)))
  rows=c.execute(select(db.outbox.c.id,db.outbox.c.payload,db.outbox.c.created).where(db.outbox.c.org_id==a['org_id'],db.outbox.c.event_type=='sale.celebrate',db.outbox.c.created>max(since,time.time()-300)).order_by(db.outbox.c.created).limit(50)).mappings().all()
  return {'items':[dict(r) for r in rows],'server_time':time.time()}

@app.get('/api/clocks')
def clocks():
 now=datetime.now(timezone.utc);pak=now.astimezone(ZoneInfo('Asia/Karachi'));zones=['America/New_York','America/Chicago','America/Denver','America/Phoenix','America/Los_Angeles','America/Anchorage','Pacific/Honolulu','Asia/Karachi']
 return {'server_time':now.isoformat(),'business_date':now.astimezone(PACIFIC).date().isoformat(),'items':[{'zone':z,'time':now.astimezone(ZoneInfo(z)).isoformat(),'abbreviation':now.astimezone(ZoneInfo(z)).tzname(),'pakistan_difference_hours':(pak.utcoffset()-now.astimezone(ZoneInfo(z)).utcoffset()).total_seconds()/3600} for z in zones]}

@app.get('/api/teams')
def list_teams(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);sec.require(a);rows=c.execute(select(db.teams).where(db.teams.c.org_id==a['org_id'],sec.scope(c,a,db.teams.c.supervisor_id))).mappings().all()
  return {'items':[{**dict(r),'member_ids':list(c.execute(select(db.members.c.user_id).where(db.members.c.team_id==r['id'])).scalars())} for r in rows]}
@app.post('/api/teams')
def create_team(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a);supervisor=payload.get('supervisor_id') or a['id'];members=list(dict.fromkeys(payload.get('member_ids',[])));leader=payload.get('leader_id') or None
  if not isinstance(members,list) or len(members)>100:raise ValueError('Use up to 100 team members.')
  for id in [supervisor,*members]:
   if not sec.allowed_user(c,a,id):raise HTTPException(403,'Team member is outside your scope.')
  sup=c.execute(select(db.users).where(db.users.c.id==supervisor)).mappings().one();sec.require(sup)
  if leader and leader not in members:raise ValueError('Group leader must be a team member.')
  row=dict(id=db.uid(),org_id=a['org_id'],name=clean(payload.get('name'),'Team name',True,150),supervisor_id=supervisor,leader_id=leader,created=time.time());c.execute(db.teams.insert().values(**row))
  if members:c.execute(db.members.insert(),[{'team_id':row['id'],'user_id':id,'created':time.time()} for id in members])
  db.log(c,a,'team.created',row['id'],{'member_count':len(members)});return {**row,'member_ids':members}
@app.get('/api/settings')
def get_settings(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);sec.require(a,3);return {**db.policy(c),'domains':[dict(r) for r in c.execute(select(db.domains)).mappings()],'projects':[dict(r) for r in c.execute(select(db.projects)).mappings()]}
@app.patch('/api/settings')
def save_settings(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a,3);old=db.policy(c);p={**old,**{k:payload[k] for k in ('import_edit_minutes','call_target','connection_target') if k in payload}}
  if not 10<=int(p['import_edit_minutes'])<=20 or not 1<=int(p['call_target'])<=1000 or not .01<=float(p['connection_target'])<=1:raise ValueError('Invalid target or import edit window.')
  p={**p,'version':old['version']+1};effective=reporting.schedule(c,p);p['targets_effective_date']=effective
  c.execute(db.settings.update().where(db.settings.c.key=='policy').values(value=p));db.log(c,a,'settings.updated',detail={'before':old,'after':p,'effect':'Targets apply from '+effective+' Pacific. Existing dates keep their target version; import window applies to new outcomes.'});return p
@app.post('/api/settings/domains')
def add_domain(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a,3);id=clean(payload.get('id') or db.uid(),'Domain ID',True,60);name=clean(payload.get('name'),'Domain name',True,150)
  if not re.fullmatch(r'[a-zA-Z0-9_-]+',id):raise ValueError('Invalid domain ID.')
  c.execute(db.domains.insert().values(id=id,name=name,active=True));db.log(c,a,'domain.created',id);return {'id':id,'name':name}
@app.post('/api/settings/agreements-link')
def link_own_agreements(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a,4);target=clean(payload.get('tfsc_user_id'),'TFSC account ID',True,32)
  if not re.fullmatch(r'[a-f0-9]{32}',target):raise ValueError('Use an existing TFSC user ID.')
  c.execute(db.users.update().where(db.users.c.id==a['id']).values(agreements_user_id=target,version=a['version']+1));db.log(c,a,'agreements.account_linked',detail={'tfsc_user_id':target});return {'ok':True}
@app.post('/api/settings/projects')
def add_project(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a,3);id=clean(payload.get('id'),'Project ID',True,32).upper()
  if not re.fullmatch(r'P[0-9]{4,12}',id):raise ValueError('Use a project ID such as P0001.')
  v=dict(id=id,name=clean(payload.get('name') or id,'Project name',True,150),domain_id=payload.get('domain_id') or None,owner_name=clean(payload.get('owner_name',''),'Owner name',False,150),owner_email=email(payload.get('owner_email','')),active=True)
  existing=c.execute(select(db.projects).where(db.projects.c.id==id)).first()
  c.execute(db.projects.update().where(db.projects.c.id==id).values(**v) if existing else db.projects.insert().values(**v));db.log(c,a,'project.updated' if existing else 'project.created',id);return v
@app.get('/api/dashboard/layout')
def get_layout(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);sec.require(a);return {'layout':c.execute(select(db.layouts.c.layout).where(db.layouts.c.user_id==a['id'])).scalar_one_or_none() or []}
@app.post('/api/dashboard/layout')
def save_layout(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a);layout=payload.get('layout',[])
  if not isinstance(layout,list) or len(layout)>20 or len(json.dumps(layout))>5000:raise ValueError('Invalid layout.')
  c.execute(delete(db.layouts).where(db.layouts.c.user_id==a['id']));c.execute(db.layouts.insert().values(user_id=a['id'],layout=layout));return {'layout':layout}

@app.get('/api/meetings')
def list_meetings(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);sec.require(a);rows,cursor=page(select(db.meetings).where(db.meetings.c.org_id==a['org_id'],sec.scope(c,a,db.meetings.c.owner_id)),db.meetings,request,c);return {'items':rows,'next_cursor':cursor}
@app.post('/api/meetings')
def create_meeting(payload:dict,request:Request):
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a);lead=get_record(c,a,db.leads,payload.get('lead_id'));start=clean(payload.get('starts_at'),'Meeting time',True,50)
  d=datetime.fromisoformat(start.replace('Z','+00:00'))
  if not d.tzinfo:raise ValueError('Meeting time requires an explicit timezone offset.')
  zone=clean(payload.get('timezone') or 'America/Los_Angeles','Timezone',True,80);ZoneInfo(zone)
  link=clean(payload.get('join_url',''),'Meeting link',False,2000)
  if link and (urlsplit(link).scheme!='https' or urlsplit(link).hostname!='meet.google.com'):raise ValueError('Use a verified HTTPS Google Meet link.')
  v=dict(id=db.uid(),org_id=a['org_id'],lead_id=lead['id'],owner_id=lead['owner_id'],title=clean(payload.get('title') or 'Client meeting','Title',True,200),starts_at=d.astimezone(timezone.utc).isoformat(),ends_at=None,timezone=zone,join_url=link or None,provider_id=None,status='manually_scheduled',created=time.time())
  c.execute(db.meetings.insert().values(**v));db.log(c,a,'meeting.recorded',v['id'],{'source':'manual'},lead['owner_id']);return v
@app.get('/api/integrations/status')
def integrations(request:Request):
 with db.engine.connect() as c:
  a=sec.actor(request,c);sec.require(a)
  return {'agreements':{'configured':len(os.getenv('CRM_TFSC_SECRET',''))>=32,'entitled':bool(a['agreements_user_id'])},'google_calendar':{'configured':google_calendar.configured(),'connected':bool(c.execute(select(google_calendar.connections.c.user_id).where(google_calendar.connections.c.user_id==a['id'])).first())},'vonage':{'configured':False,'status':'Uses the registered system telephone handler. Register your Vonage desktop client to handle tel: links.'},'email':{'configured':False},'esign':{'configured':False}}
@app.post('/api/leads/{id}/agreements')
def handoff(id:str,payload:dict,request:Request):
 secret=os.getenv('CRM_TFSC_SECRET');url=os.getenv('CRM_TFSC_URL','http://127.0.0.1:8765').rstrip('/')
 if not secret:raise HTTPException(503,'The agreements bridge needs administrator configuration.')
 with db.engine.begin() as c:
  a=sec.actor(request,c);sec.require(a)
  if not a['agreements_user_id']:raise HTTPException(403,'Your account has no agreements entitlement.')
  lead=get_record(c,a,db.leads,id)
  if lead['stage']!='signed_up':raise HTTPException(409,'Confirm the sale before proceeding to agreements.')
  activity=c.execute(select(db.activities).where(db.activities.c.id==lead['activity_id'])).mappings().one()
  entity=payload.get('entity')
  if entity not in ('SCRIBE ALIGN LLC','EXTENDED HEALTH SERVICES INC.','24-7 CONSULTANCY LLC'):raise ValueError('Choose a supported legal entity.')
  key=clean(payload.get('idempotency_key'),'Handoff request key',True,80)
  body={'idempotency_key':key,'crm_client_id':id,'crm_actor_id':a['id'],'entity':entity,'client':{'client_name':payload.get('client_name') or activity['provider_name'],'address':payload.get('address') or activity['address'],'client_email':payload.get('client_email') or activity['email'],'client_signatory':payload.get('client_signatory',''),'client_title':payload.get('client_title','')},'services':payload.get('services',[]),'document_types':payload.get('document_types',['baa']),'effective_date':iso_date(payload.get('effective_date'))}
  existing=c.execute(select(db.handoffs).where(db.handoffs.c.owner_id==a['id'],db.handoffs.c.key==key)).mappings().first()
  if existing:
   if existing['body_hash']!=body_hash(body):raise HTTPException(409,'Handoff key already used for different details.')
   # Replay through TFSC to obtain a fresh, single-use redemption code.
 # Do not keep a DB transaction open while waiting for the other site.
 raw=json.dumps(body,separators=(',',':')).encode();timestamp=str(int(time.time()));nonce=secrets.token_urlsafe(24);path='/api/integrations/v1/crm/handoffs'
 canonical='\n'.join([timestamp,nonce,'POST',path,hashlib.sha256(raw).hexdigest()]);signature=hmac.new(secret.encode(),canonical.encode(),hashlib.sha256).hexdigest()
 try:
  response=httpx.post(os.getenv('CRM_TFSC_INTERNAL_URL',url).rstrip('/')+path,content=raw,headers={'Host':urlsplit(url).netloc,'Content-Type':'application/json','X-TFSC-Timestamp':timestamp,'X-TFSC-Nonce':nonce,'X-TFSC-Signature':signature},timeout=20)
  if response.status_code>=400:raise HTTPException(502,'TFSC could not accept the handoff. Verify the account mapping and selected company in TFSC.')
  result=response.json()
 except httpx.HTTPError:raise HTTPException(503,'TFSC is unavailable. Retry with the same request key.')
 with db.engine.begin() as c:
  current=sec.actor(request,c);sec.require(current)
  if not current['agreements_user_id']:raise HTTPException(403,'Agreements access has been revoked.')
  saved={k:v for k,v in result.items() if k!='redeem_token'}
  if existing:c.execute(db.handoffs.update().where(db.handoffs.c.id==existing['id']).values(result=saved))
  else:c.execute(db.handoffs.insert().values(id=db.uid(),org_id=a['org_id'],lead_id=id,owner_id=a['id'],key=key,body_hash=body_hash(body),result=saved,created=time.time()))
  db.log(c,a,'agreement.handoff_created',id,{'entity':entity},lead['owner_id'])
 return {**result,'agreements_url':url+'/reviews'}

followups.install(app,clean)
import crm_contracts
crm_contracts.install(app,get_record)
import google_calendar
google_calendar.install(app,ORIGIN,get_record,clean)
import oidc_support
oidc_support.install(app,db,sec,session_response,ORIGIN,'CRM',sec.actor)
WEB=Path(os.getenv('CRM_WEB_DIST') or (db.APP/'web'/'dist' if (db.APP/'web'/'dist').exists() else db.APP.parent/'crm-web'/'dist'))
@app.get('/')
@app.get('/login')
def index():
 path=WEB/'index.html'
 if not path.exists():return JSONResponse({'app':'TFSC Sales CRM','status':'Frontend build required. Run npm install and npm run build in web.'},503)
 return FileResponse(path)
@app.get('/assets/{path:path}')
def asset(path:str):
 file=(WEB/'assets'/path).resolve();root=(WEB/'assets').resolve()
 if not file.is_relative_to(root) or not file.is_file():raise HTTPException(404,'Asset not found.')
 return FileResponse(file)
@app.get('/{path:path}')
def public_file(path:str):
 file=(WEB/path).resolve()
 if file.is_relative_to(WEB.resolve()) and file.is_file() and file.suffix in ('.png','.svg','.ico','.woff2'):return FileResponse(file)
 raise HTTPException(404,'Page not found.')
