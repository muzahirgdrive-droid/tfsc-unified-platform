"""Legacy tracking-workbook reader and reconciliation (§20.5, §25.3). Pure: no database.

Cell contents are data, never instructions. Formulas are never imported as authority: formula cells in the
input columns (A–O) are treated as empty, and the workbook's own score/stage columns (P/Q) are only inspected
for broken references (#REF!), which go to the review queue.

Workbook rules reproduced for the reconciliation figures (from the sheet's own formulas, §2):
- a "Yes" cell is a case-insensitive match of the text "Yes" (Excel `E2="Yes"` semantics, no trimming);
- row score = 100 sign-up · 50 strong FU · 30 connected + (text or email) · 20 connected · 10 attempted,
  only when Call Attempt = Yes; quality = Σ score ÷ (100 × attempts).
"""
from __future__ import annotations

import hashlib
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

# Workbook columns A–O (P = Quality %, Q = Stage are formulas and never read as data).
COLUMNS = ('business_date', 'name', 'phone', 'email', 'called', 'connected', 'text_sent', 'email_sent', 'strong_followup',
           'sign_up', 'followup_count', 'last_followup', 'next_followup', 'call_again', 'remarks')
CALLED, CONNECTED, TEXT, EMAIL, STRONG, SIGNUP = 4, 5, 6, 7, 8, 9  # workbook columns E–J
NON_OPERATOR = {'dashboard', 'validations', 'o'}
MAX_ROWS_PER_SHEET = 200_000

# §2 / §20.5: raw-input figures of the audited workbook for 1–14 Oct 2026.
AUDITED_SHA256 = '278fe1439b6d58e9eca011493cdce62fba74f354e762dbe4b3abf70b75c93f2d'
AUDITED_PERIOD = (date(2026, 10, 1), date(2026, 10, 14))
AUDITED_FIGURES = {'attempts': 17020, 'connected': 2803, 'conn_rate': 0.164689, 'signups': 1, 'strong_fu': 27,
                   'quality': 0.118143, 'contradictions': 33, 'ref_cells': 21}


def is_yes(v: Any) -> bool:
    return isinstance(v, str) and v.lower() == 'yes'


def sheet_skip_reason(title: str, state: str) -> str | None:
    t = title.strip().lower()
    if t in NON_OPERATOR:
        return 'not an operator sheet'
    if t.startswith('unassigned'):
        return 'unassigned slot'
    if state != 'visible':
        return 'hidden sheet'
    return None


@dataclass
class Row:
    sheet: str
    row_no: int  # Excel row number
    values: tuple  # 15 input cells A–O (formula cells -> None)
    ref_cells: list[str] = field(default_factory=list)

    @property
    def business_date(self) -> date | None:
        return to_date(self.values[0])

    @property
    def contradiction(self) -> bool:
        """An outcome is Yes while Call Attempt is not (defect #5, §2)."""
        return not is_yes(self.values[CALLED]) and any(is_yes(self.values[i]) for i in (CONNECTED, TEXT, EMAIL, STRONG, SIGNUP))

    def points(self) -> int:
        v = self.values
        if not is_yes(v[CALLED]):
            return 0
        if is_yes(v[SIGNUP]):
            return 100
        if is_yes(v[STRONG]):
            return 50
        if is_yes(v[CONNECTED]) and (is_yes(v[TEXT]) or is_yes(v[EMAIL])):
            return 30
        return 20 if is_yes(v[CONNECTED]) else 10

    def review_codes(self) -> list[str]:
        codes = []
        if self.contradiction:
            codes.append('contradiction')
        if self.business_date is None:
            codes.append('missing_date')
        if self.ref_cells:
            codes.append('formula_error')
        return codes

    @property
    def importable(self) -> bool:
        """Contradictions and undated rows wait in the review queue; #REF! rows import (their inputs are intact)."""
        return not self.contradiction and self.business_date is not None

    def as_input(self) -> dict[str, Any]:
        """The row in the shape the refine engine / activities.create expect (strings, ISO dates)."""
        out: dict[str, Any] = {}
        for k, v in zip(COLUMNS, self.values, strict=True):
            if isinstance(v, datetime):
                v = v.date().isoformat()
            elif isinstance(v, date):
                v = v.isoformat()
            elif isinstance(v, float) and v.is_integer():
                v = str(int(v))
            elif v is not None and not isinstance(v, str):
                v = str(v)
            out[k] = v
        return out


def to_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ('%m/%d/%Y', '%Y-%m-%d', '%m/%d/%y', '%m-%d-%Y'):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


@dataclass
class Scan:
    sha256: str
    sheets: list[dict]
    roster: list[str]
    rows: list[Row]

    def operator_sheets(self) -> list[str]:
        return [s['name'] for s in self.sheets if not s.get('skipped')]

    def summary(self) -> dict:
        return {'sha256': self.sha256, 'sheets': self.sheets, 'roster': self.roster, 'rows': len(self.rows),
                'attempts_all_time': sum(1 for r in self.rows if is_yes(r.values[CALLED])),
                'review': dict(Counter(code for r in self.rows for code in r.review_codes())),
                'ref_cells': sum(len(r.ref_cells) for r in self.rows), 'contradictions': sum(1 for r in self.rows if r.contradiction)}


def _cell_value(c: Any) -> Any:
    if c.data_type == 'f':
        return None
    v = c.value
    return None if v == '' else v  # Excel blank semantics: a cell holding a space is not blank (COUNTA counts it)


def scan(path: str | Path) -> Scan:
    """Read every operator sheet. Rows count as entries when they carry a provider name or phone."""
    import openpyxl  # defusedxml is installed, so openpyxl parses the XML defensively

    p = Path(path)
    digest = hashlib.sha256(p.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(p, read_only=True, data_only=False)
    try:
        sheets: list[dict] = []
        rows: list[Row] = []
        roster: list[str] = []
        for ws in wb.worksheets:
            state = getattr(ws, 'sheet_state', 'visible')
            if ws.title.strip().lower() == 'validations':
                roster = _roster(ws)
            reason = sheet_skip_reason(ws.title, state)
            info: dict[str, Any] = {'name': ws.title, 'state': state}
            if reason:
                sheets.append({**info, 'skipped': reason})
                continue
            stats: Counter[str] = Counter()
            for n, cells in enumerate(ws.iter_rows(min_row=2, max_col=17), start=2):
                if n > MAX_ROWS_PER_SHEET + 1:
                    stats['truncated'] += 1
                    break
                cells = tuple(cells) + (None,) * (17 - len(cells))
                inputs = tuple(_cell_value(c) if c is not None else None for c in cells[:15])
                refs = [c.coordinate for c in cells if c is not None and _is_ref_error(c)]
                if inputs[1] is None and inputs[2] is None:
                    if refs:
                        stats['formula_error_without_entry'] += 1
                    if any(v is not None for v in inputs):
                        stats['no_name_or_phone'] += 1  # reported, not imported: an entry needs a provider name or phone
                        stats['no_name_or_phone_attempts'] += int(is_yes(inputs[CALLED]))
                    elif any(c is not None and c.data_type == 'f' for c in cells):
                        stats['formula_only'] += 1
                    continue
                rows.append(Row(ws.title, n, inputs, refs))
                stats['entries'] += 1
            sheets.append({**info, **stats})
        return Scan(digest, sheets, roster, rows)
    finally:
        wb.close()


def _is_ref_error(c: Any) -> bool:
    v = c.value
    return isinstance(v, str) and '#REF!' in v and c.data_type in ('f', 'e')


def _roster(ws: Any) -> list[str]:
    """Validations!C32:E54 lists the reps (names only; names are not identities)."""
    out = []
    for row in ws.iter_rows(min_row=32, max_row=54, min_col=3, max_col=5, values_only=True):
        name = row[0]
        if isinstance(name, str) and name.strip():
            out.append(name.strip()[:120])
    return out


def default_period(s: Scan) -> tuple[date, date]:
    """The audited workbook reconciles 1–14 Oct 2026 (§20.5); any other workbook its full date span."""
    if s.sha256 == AUDITED_SHA256:
        return AUDITED_PERIOD
    dates = [d for r in s.rows if (d := r.business_date) is not None]
    return (min(dates), max(dates)) if dates else AUDITED_PERIOD


def figures(rows: list[Row], start: date, end: date, *, importable_only: bool = False) -> dict[str, Any]:
    """Raw-input figures for a period, using the workbook's rules (not its cached formulas)."""
    sel = [r for r in rows if (d := r.business_date) is not None and start <= d <= end and (r.importable or not importable_only)]
    att = [r for r in sel if is_yes(r.values[CALLED])]
    attempts = len(att)
    points = sum(r.points() for r in att)
    connected = sum(1 for r in sel if is_yes(r.values[CONNECTED]))
    return {
        'rows': len(sel), 'attempts': attempts, 'connected': connected,
        'conn_rate': round(connected / attempts, 6) if attempts else 0.0,
        'sms': sum(1 for r in att if is_yes(r.values[TEXT])), 'emails': sum(1 for r in att if is_yes(r.values[EMAIL])),
        'strong_fu': sum(1 for r in sel if is_yes(r.values[STRONG])), 'signups': sum(1 for r in sel if is_yes(r.values[SIGNUP])),
        'points': points, 'quality': round(points / (100 * attempts), 6) if attempts else 0.0,
    }


def phone_repeats(rows: list[Row]) -> dict[str, int]:
    """Repeated 10-digit phone keys (defect #7): keys seen more than once, and keys seen on different sheets.
    Key = the phone's digits, a leading US country code 1 removed, exactly 10 digits."""
    by_key: dict[str, set[str]] = {}
    counts: Counter[str] = Counter()
    for r in rows:
        digits = ''.join(ch for ch in str(r.values[2] or '') if ch.isdigit())
        if len(digits) == 11 and digits.startswith('1'):
            digits = digits[1:]
        if len(digits) != 10:
            continue
        counts[digits] += 1
        by_key.setdefault(digits, set()).add(r.sheet)
    return {'repeated_keys': sum(1 for n in counts.values() if n > 1),
            'cross_operator_keys': sum(1 for s in by_key.values() if len(s) > 1)}


def reconciliation(s: Scan, start: date, end: date, database: dict | None = None, sheets: list[str] | None = None) -> dict[str, Any]:
    """Workbook figures vs (optionally) what the CRM holds after import, plus the audited §2 figures when this is
    the audited workbook. `importable` (rows the import writes, limited to the mapped `sheets`) is what the
    database must match exactly."""
    wb = figures(s.rows, start, end)
    wb['contradictions'] = sum(1 for r in s.rows if r.contradiction)
    wb['ref_cells'] = sum(len(r.ref_cells) for r in s.rows)
    out: dict[str, Any] = {'period': {'start': start.isoformat(), 'end': end.isoformat()}, 'workbook': wb,
                           'importable': figures([r for r in s.rows if sheets is None or r.sheet in sheets], start, end, importable_only=True),
                           'phone_repeats': phone_repeats(s.rows)}
    if s.sha256 == AUDITED_SHA256 and (start, end) == AUDITED_PERIOD:
        out['expected'] = AUDITED_FIGURES
        out['workbook_matches_expected'] = {k: _eq(wb.get(k), v) for k, v in AUDITED_FIGURES.items()}
    if database is not None:
        out['database'] = database
        ref = out['importable']
        out['database_matches_workbook'] = {k: _eq(database.get(k), ref[k]) for k in ('attempts', 'connected', 'conn_rate', 'strong_fu',
                                                                                     'signups', 'points', 'quality') if k in database}
    return out


def _eq(a: Any, b: Any) -> bool:
    if isinstance(a, float) or isinstance(b, float):
        return a is not None and b is not None and abs(float(a) - float(b)) < 5e-7
    return a == b
