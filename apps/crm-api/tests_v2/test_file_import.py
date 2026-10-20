"""XLSX/CSV import job (§11.2): header detection, per-row results, live rules, idempotent re-import, resume, limits."""
import io
from datetime import timedelta

from sqlalchemy import text

from .conftest import new_phone


def _csv(rows: list[list[str]]) -> bytes:
    return ('\n'.join(','.join(f'"{c}"' for c in r) for r in rows) + '\n').encode()


def _upload(client, data: bytes, name: str):
    return client.post('/api/imports', files={'file': (name, io.BytesIO(data), 'application/octet-stream')})


def test_csv_import_saves_valid_rows_and_reports_the_rest(ws):
    b = ws['b']
    p1, p2, p3, p5 = new_phone(), new_phone(), new_phone(), new_phone()
    data = _csv([['Provider Name', 'Phone', 'Call Attempt', 'Connected', 'Remarks'],
                 ['sunrise pediatrics pa', p1, 'Yes', 'No', 'left voicemail'],
                 ['harbor family medicine', p2, 'Y', 'Y', ''],
                 ['metro imaging llc', p3, '', '', 'research only'],
                 ['bad number clinic', '12345', 'Yes', 'No', ''],
                 ['contradiction care', p5, 'No', 'Yes', '=HYPERLINK("http://evil.example","click")']])
    r = _upload(b, data, 'day.csv')
    assert r.status_code == 201, r.text
    up = r.json()
    assert up['sheets'] == [{'name': 'CSV', 'rows': 5, 'too_many': False, 'header_detected': True,
                             'mapping': ['name', 'phone', 'called', 'connected', 'remarks'],
                             'headers': ['Provider Name', 'Phone', 'Call Attempt', 'Connected', 'Remarks']}]
    r = b.post(f"/api/imports/{up['id']}/start", json={'sheet': 'CSV'})
    assert r.status_code == 200, r.text
    job = r.json()
    assert job['status'] == 'done' and job['result'] == {'saved': 3, 'rejected': 1, 'review': 1}
    assert job['progress'] == {'done': 5, 'total': 5}
    rows = b.get(f"/api/imports/{up['id']}/rows", params={'decision': 'rejected'}).json()['items']
    assert [r['row_no'] for r in rows] == [3] and rows[0]['issues'][0]['field'] == 'phone'
    review = b.get(f"/api/imports/{up['id']}/rows", params={'decision': 'review'}).json()['items']
    assert [r['row_no'] for r in review] == [4]
    # Live rules: source=import, server Pacific date, the import lock window (10–20 min, default 15).
    from tfsc_crm.db.session import system_tx
    with system_tx() as c:
        acts = c.execute(text('SELECT entry_source, first_committed_at, edit_deadline, remarks FROM activities WHERE import_batch_id = :b ORDER BY created_at'),
                         {'b': up['id']}).all()
    assert len(acts) == 3 and {a[0] for a in acts} == {'import'}
    committed = [a for a in acts if a[1] is not None]
    assert len(committed) == 2 and all(a[2] - a[1] == timedelta(minutes=15) for a in committed)
    assert next(a for a in acts if a[1] is None)[2] is None  # data hunting: no timer


def test_reimporting_the_same_sheet_is_refused_and_others_cannot_see_it(ws):
    b = ws['b']
    data = _csv([['Provider Name', 'Phone', 'Call Attempt'], ['once only clinic', new_phone(), 'Yes']])
    first = _upload(b, data, 'once.csv').json()
    assert b.post(f"/api/imports/{first['id']}/start", json={'sheet': 'CSV'}).json()['status'] == 'done'
    again = _upload(b, data, 'once-again.csv').json()
    r = b.post(f"/api/imports/{again['id']}/start", json={'sheet': 'CSV'})
    assert r.status_code == 409 and first['id'] in r.json()['error']
    assert b.post(f"/api/imports/{first['id']}/start", json={'sheet': 'CSV'}).status_code == 409
    assert ws['c'].get(f"/api/imports/{first['id']}").status_code == 404
    assert ws['c'].get(f"/api/imports/{first['id']}/rows").status_code == 404
    assert all(i['id'] != first['id'] for i in ws['c'].get('/api/imports').json())


def test_xlsx_choose_sheet_and_map_columns_without_header(ws, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = 'Notes'
    wb.active['A1'] = 'not data'
    s2 = wb.create_sheet('Monday')
    s2.append(['alpha ortho & spine group', new_phone(), 'Yes', 'Yes', 'Yes'])
    s2.append(['bay area aba therapy inc', new_phone(), 'Yes', 'No', 'No'])
    hidden = wb.create_sheet('Hidden')
    hidden.sheet_state = 'hidden'
    path = tmp_path / 'week.xlsx'
    wb.save(path)
    a = ws['a']
    up = _upload(a, path.read_bytes(), 'week.xlsx').json()
    assert [s['name'] for s in up['sheets']] == ['Notes', 'Monday']  # hidden sheets are not offered
    mon = up['sheets'][1]
    assert mon['header_detected'] is False and mon['rows'] == 2
    r = a.post(f"/api/imports/{up['id']}/start", json={'sheet': 'Monday', 'columns': ['name', 'phone', 'called', 'connected', 'text_sent']})
    assert r.json()['result'] == {'saved': 2, 'rejected': 0, 'review': 0}
    assert a.post(f"/api/imports/{up['id']}/start", json={'sheet': 'Notes'}).status_code == 409
    up2 = _upload(a, path.read_bytes(), 'week.xlsx').json()
    assert a.post(f"/api/imports/{up2['id']}/start", json={'sheet': 'Notes', 'columns': ['remarks']}).status_code == 422


def test_import_for_someone_in_scope_only(ws):
    data = _csv([['Provider Name', 'Phone', 'Call Attempt'], ['scope clinic', new_phone(), 'Yes']])
    up = _upload(ws['sup'], data, 'scope.csv').json()
    r = ws['sup'].post(f"/api/imports/{up['id']}/start", json={'sheet': 'CSV', 'owner_id': ws['c'].id})  # c reports to sup2
    assert r.status_code == 403
    r = ws['sup'].post(f"/api/imports/{up['id']}/start", json={'sheet': 'CSV', 'owner_id': ws['b'].id})
    assert r.status_code == 200 and r.json()['result']['saved'] == 1
    up = _upload(ws['a'], _csv([['Provider Name', 'Phone'], ['x', new_phone()]]), 'x.csv').json()
    assert ws['a'].post(f"/api/imports/{up['id']}/start", json={'sheet': 'CSV', 'owner_id': ws['b'].id}).status_code == 403


def test_resume_after_a_crash_does_not_duplicate(ws, monkeypatch):
    from tfsc_crm.domain import file_import as fi
    b = ws['b']
    data = _csv([['Provider Name', 'Phone', 'Call Attempt']] + [[f'resume clinic {i}', new_phone(), 'Yes'] for i in range(7)])
    up = _upload(b, data, 'resume.csv').json()
    monkeypatch.setattr(fi, 'CHUNK', 2)
    real, calls = fi.process_row, {'n': 0}

    def flaky(*args, **kw):
        calls['n'] += 1
        if calls['n'] == 5:
            raise RuntimeError('worker lost')
        return real(*args, **kw)

    monkeypatch.setattr(fi, 'process_row', flaky)
    job = b.post(f"/api/imports/{up['id']}/start", json={'sheet': 'CSV'}).json()
    assert job['status'] == 'failed' and job['progress']['done'] == 4
    job = b.post(f"/api/imports/{up['id']}/start", json={'sheet': 'ignored on resume'}).json()
    assert job['status'] == 'done' and job['result']['saved'] == 7
    from tfsc_crm.db.session import system_tx
    with system_tx() as c:
        assert c.execute(text('SELECT count(*) FROM activities WHERE import_batch_id = :b'), {'b': up['id']}).scalar_one() == 7


def test_limits(ws):
    a = ws['a']
    assert _upload(a, b'hello', 'notes.pdf').status_code == 400
    big = _csv([['Provider Name', 'Phone']] + [[f'p{i}', '4155550100'] for i in range(50_001)])
    up = _upload(a, big, 'big.csv').json()
    assert up['sheets'][0]['too_many'] is True
    assert a.post(f"/api/imports/{up['id']}/start", json={'sheet': 'CSV'}).status_code == 422
