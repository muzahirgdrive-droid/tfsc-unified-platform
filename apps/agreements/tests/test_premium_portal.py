"""Version 5: company settings and signatures, the completion rule, links that do not expire, smart field placement, the
client document library, Home figures, reports (screen, Excel, CSV), the activity log, and who may use each of them."""
import base64
import io
import sys
import time
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from sqlalchemy import func, select, update

import companies
import database as db
import esign
import esign_db as t
import esign_pdf
import server as legacy
from document_engine import generate_docx
from esign_support import agreement_pdf, signature_png
from test_esign import EHS, SCRIBE, Base


def photo_of_signature() -> bytes:
    """A JPEG like a phone photo of a signature: grey paper, blue ink."""
    image = Image.new('RGB', (900, 500), (214, 211, 205))
    draw = ImageDraw.Draw(image)
    draw.line([(150, 300), (300, 180), (420, 320), (560, 170), (720, 300)], fill=(30, 40, 120), width=9)
    out = io.BytesIO()
    image.save(out, 'JPEG', quality=88)
    return out.getvalue()


class CompanySettings(Base):
    def test_only_the_super_admin_manages_company_settings(self):
        image = base64.b64encode(photo_of_signature()).decode()
        for user in (self.admin, self.u1):
            self.login(user)
            self.assertEqual(self.api.get('/api/admin/companies').status_code, 403)
            self.assertEqual(self.post('/api/admin/companies/scribe', {'signatory': 'Someone Else'}).status_code, 403)
            self.assertEqual(self.post('/api/admin/companies/clean-signature', {'image': image}).status_code, 403)
            self.assertEqual(self.post('/api/admin/workspace', {'help_email': 'x@example.com'}).status_code, 403)
            self.assertEqual(self.post('/api/admin/email-trust', {'company': 'scribe'}).status_code, 403)
            self.assertEqual(self.post('/api/signature', {'entity': SCRIBE, 'image': image}).status_code, 403)
        self.login(self.super)
        self.assertEqual(self.post('/api/admin/companies/nope', {'signatory': 'X'}).status_code, 404)
        r = self.post('/api/admin/companies/scribe', {'signatory': 'Naeem Abbas', 'title': 'Chief Executive Officer', 'signature': image, 'signature_kind': 'uploaded'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(r.json()['company']['has_signature'])
        self.assertEqual(r.json()['company']['signature_source'], 'settings')

        # The image is served only to people who work for that company.
        self.login(self.ehs_user)
        self.assertEqual(self.api.get('/api/companies/scribe/signature.png').status_code, 403)
        self.login(self.u1)
        png = self.api.get('/api/companies/scribe/signature.png')
        self.assertEqual(png.status_code, 200)
        with Image.open(io.BytesIO(png.content)) as cleaned:
            self.assertEqual(cleaned.mode, 'RGBA')
            self.assertEqual(cleaned.getpixel((0, 0))[3], 0, 'the paper is transparent')
            self.assertLess(cleaned.width, 700, 'trimmed to the ink')
        config = self.api.get('/api/config').json()
        scribe = next(e for e in config['entities'] if e['name'] == SCRIBE)
        self.assertTrue(scribe['signature_available'])
        self.assertEqual(config['workspace']['help_email'], 'muzahir@scribealign.com')

    def test_signatory_and_title_and_signature_reach_the_agreement(self):
        self.login(self.super)
        r = self.post('/api/admin/companies/24-7', {'signatory': 'Danish Ather', 'title': 'Chief Operating Officer',
                                                    'signature': signature_png('Danish Ather', (600, 160)), 'signature_kind': 'typed'})
        self.assertEqual(r.status_code, 200, r.text)
        with db.engine.connect() as c:
            company = companies.profile(c, '24-7 CONSULTANCY LLC')
        self.assertTrue(company['png'].startswith(b'\x89PNG'))
        payload = legacy.validate({'entity': '24-7 CONSULTANCY LLC', 'client_name': 'Northside Clinic', 'address': '1 Main St, Atlanta, GA 30301',
                                   'effective_date': '2026-10-24', 'document_types': ['nda'], 'output_root': str(self.root)})
        docs = legacy.build_documents(payload, 'TEST-SIG', dict(company, title='Chief Operating Officer'))
        target = self.root / 'nda.docx'
        generate_docx(docs[0]['blocks'], docs[0]['data'], 'TEST-SIG', target)
        from zipfile import ZipFile
        with ZipFile(target) as z:
            self.assertIn('word/media/sa-signature.png', z.namelist())
            xml = z.read('word/document.xml').decode()
        self.assertIn('Name: Danish Ather', xml)
        self.assertIn('Title: Chief Operating Officer', xml)

    def test_signature_size_defaults_and_setting_reach_the_agreement(self):
        from zipfile import ZipFile
        import re as re_
        self.login(self.super)
        sizes = {x['key']: (x['signature_scale'], x['default_signature_scale']) for x in self.api.get('/api/admin/companies').json()['companies']}
        self.assertEqual(sizes, {'scribe': (0.6, 0.6), 'ehs': (0.6, 0.6), '24-7': (0.85, 0.85)}, 'Naeem Abbas 60 %, Danish Ather 85 %')
        for bad in (1.5, 0.1, 'big', True):
            self.assertEqual(self.post('/api/admin/companies/scribe', {'signature_scale': bad}).status_code, 400, bad)
        r = self.post('/api/admin/companies/scribe', {'signature': signature_png('Naeem Abbas', (600, 400)), 'signature_kind': 'drawn'})
        self.assertEqual(r.status_code, 200, r.text)

        def printed_height() -> int:
            with db.engine.connect() as c:
                company = companies.profile(c, SCRIBE)
            payload = legacy.validate({'entity': SCRIBE, 'client_name': 'Northside Clinic', 'address': '1 Main St, Reno, NV 89501',
                                       'effective_date': '2026-10-24', 'document_types': ['baa'], 'output_root': str(self.root)})
            docs = legacy.build_documents(payload, 'TEST-SIZE', company)
            target = self.root / 'size.docx'
            generate_docx(docs[0]['blocks'], docs[0]['data'], 'TEST-SIZE', target)
            with ZipFile(target) as z:
                return int(re_.search(r'<wp:extent cx="\d+" cy="(\d+)"', z.read('word/document.xml').decode()).group(1))
        from signature_store import MAX_HEIGHT
        self.assertAlmostEqual(printed_height(), MAX_HEIGHT * 0.6, delta=2, msg='the default size: 60 % of the largest box')
        r = self.post('/api/admin/companies/scribe', {'signature_scale': 0.75})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['company']['signature_scale'], 0.75)
        self.assertAlmostEqual(printed_height(), MAX_HEIGHT * 0.75, delta=2)
        self.post('/api/admin/companies/scribe', {'signature_scale': 0.6})
        with db.engine.connect() as c:
            stored = c.execute(select(db.company_profiles.c.signature_scale).where(db.company_profiles.c.entity == SCRIBE)).scalar()
        self.assertIsNone(stored, 'choosing the default again stores nothing, so it follows the default')
        self.login(self.admin)
        self.assertEqual(self.post('/api/admin/companies/scribe', {'signature_scale': 0.9}).status_code, 403)

    def test_rejects_images_without_a_signature(self):
        blank = Image.new('RGB', (400, 200), (250, 250, 250))
        out = io.BytesIO()
        blank.save(out, 'PNG')
        with self.assertRaises(ValueError):
            companies.clean_signature(out.getvalue())
        with self.assertRaises(ValueError):
            companies.clean_signature(b'%PDF-1.7 not an image')
        self.login(self.super)
        self.assertEqual(self.post('/api/admin/companies/scribe', {'signature': 'not-base64!', 'signature_kind': 'uploaded'}).status_code, 400)

    def test_workspace_settings(self):
        self.assertEqual(self.api.get('/api/public/workspace').json()['help_email'], 'muzahir@scribealign.com')
        self.login(self.super)
        self.assertEqual(self.post('/api/admin/workspace', {'help_email': 'not an email'}).status_code, 400)
        r = self.post('/api/admin/workspace', {'help_email': 'help@scribealign.com', 'copyright_holder': 'TFSC'})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.api.get('/api/public/workspace').json()['copyright_holder'], 'TFSC')
        # Emails name the Help address too
        import mailer
        from entities import get_entity
        self.assertIn('help@scribealign.com', mailer.layout(get_entity(SCRIBE), 'Please sign', '<p>Hello</p>', security_note=True))
        self.assertIn('help@scribealign.com', mailer.text_footer(get_entity(SCRIBE), None, True))
        # Every page without an account shows the copyright and the Help address (filled in from these settings)
        for page in ('/login', '/verify'):
            html = self.api.get(page).text
            self.assertIn('muzahir@scribealign.com', html, page)
            self.assertIn('All rights reserved.', html, page)

    def test_email_trust_reads_dns(self):
        answers = {'example.com': ['v=spf1 include:_spf.google.com ~all'], 'google._domainkey.example.com': ['v=DKIM1; k=rsa; p=MIIB'],
                   '_dmarc.example.com': ['v=DMARC1; p=none'], 'default._bimi.example.com': []}
        with patch.object(companies, 'txt_records', side_effect=lambda name: answers.get(name, [])):
            result = companies.email_trust('example.com', 'google', 'smtp.gmail.com')
        states = {x['key']: x['state'] for x in result['checks']}
        self.assertEqual(states, {'spf': 'ok', 'dkim': 'ok', 'dmarc': 'warning', 'bimi': 'missing'})
        self.assertIn('p=quarantine', next(x for x in result['checks'] if x['key'] == 'dmarc')['record']['value'])
        with self.assertRaises(ValueError):
            companies.email_trust('not a domain', 'google')


class CompletionRule(Base):
    def envelope(self, people, settings=None):
        self.connect_mailbox()
        self.login(self.u1)
        env = self.post('/api/esign/envelopes', {'entity': SCRIBE}).json()
        r = self.api.post(f"/api/esign/envelopes/{env['id']}/documents", json={'name': 'Order.pdf', 'data': base64.b64encode(agreement_pdf()).decode()},
                          headers={'X-CSRF-Token': self.csrf})
        env = r.json()
        doc = env['documents'][0]['id']
        fields = [{'recipient': p['key'], 'document_id': doc, 'page': 1, 'kind': 'signature', 'x': 60, 'y': 500 + i * 60, 'w': 170, 'h': 40}
                  for i, p in enumerate(people) if p['role'] == 'signer']
        body = {'subject': 'Please sign', 'recipients': people, 'fields': fields}
        if settings:
            body['settings'] = settings
        r = self.post(f"/api/esign/envelopes/{env['id']}/draft", body)
        self.assertEqual(r.status_code, 200, r.text)
        r = self.post(f"/api/esign/envelopes/{env['id']}/send")
        self.assertEqual(r.status_code, 200, r.text)
        self.drain()
        return r.json()['envelope']

    def sign(self, address):
        client, base, state = self.signer(self.smtp.link(self.smtp.to(address)[-1]))
        self.assertEqual(self.complete_signing(client, base, state).status_code, 200)
        self.drain()

    def test_viewers_do_not_hold_up_completion_and_still_get_the_documents(self):
        env = self.envelope([{'key': 'a', 'role': 'signer', 'routing_order': 1, 'name': 'Jane Client', 'email': 'jane@client.test'},
                             {'key': 'b', 'role': 'viewer', 'routing_order': 1, 'name': 'Vic Viewer', 'email': 'vic@client.test'}])
        self.assertIsNone(env['expires_at'], 'no expiration unless the sender sets one')
        self.sign('jane@client.test')
        self.assertEqual(self.api.get(f"/api/esign/envelopes/{env['id']}").json()['status'], 'completed')
        self.assertEqual(len(self.smtp.to('vic@client.test')), 2)        # the request to view, then the completed documents

    def test_with_a_signing_order_a_viewer_waits_for_their_turn(self):
        env = self.envelope([{'key': 'a', 'role': 'viewer', 'routing_order': 1, 'name': 'Vic Viewer', 'email': 'vic@client.test'},
                             {'key': 'b', 'role': 'signer', 'routing_order': 2, 'name': 'Jane Client', 'email': 'jane@client.test'}], {'ordered': True})
        self.assertEqual(self.smtp.to('jane@client.test'), [], 'the signer waits for the viewer')
        client, base, state = self.signer(self.smtp.link(self.smtp.to('vic@client.test')[0]))
        self.assertEqual(self.sign_post(client, base, 'consent', state).status_code, 200)
        self.assertEqual(self.sign_post(client, base, 'acknowledge', state).status_code, 200)
        self.drain()
        self.assertEqual(len(self.smtp.to('jane@client.test')), 1)
        self.sign('jane@client.test')
        self.assertEqual(self.api.get(f"/api/esign/envelopes/{env['id']}").json()['status'], 'completed')

    def test_envelopes_stuck_on_a_viewer_complete_after_the_update(self):
        env = self.envelope([{'key': 'a', 'role': 'signer', 'routing_order': 1, 'name': 'Jane Client', 'email': 'jane@client.test'},
                             {'key': 'b', 'role': 'viewer', 'routing_order': 2, 'name': 'Vic Viewer', 'email': 'vic@client.test'}])
        # As 4.2 left it: the signer has signed, the viewer has not viewed, the envelope is still waiting.
        with db.engine.begin() as c:
            c.execute(update(t.recipients).where(t.recipients.c.envelope_id == env['id'], t.recipients.c.role == 'signer').values(status='signed', signed_at=time.time()))
        self.assertEqual(esign.settle(), 1)
        self.drain()
        self.assertEqual(self.api.get(f"/api/esign/envelopes/{env['id']}").json()['status'], 'completed')
        self.assertEqual(esign.settle(), 0)

    def test_expiration_only_when_the_sender_sets_it_and_links_stay_active(self):
        env = self.envelope([{'key': 'a', 'role': 'signer', 'routing_order': 1, 'name': 'Jane Client', 'email': 'jane@client.test'}], {'expire_days': 10})
        self.assertGreater(env['expires_at'], time.time() + 9 * 86400)
        r = self.post(f"/api/esign/envelopes/{env['id']}/correct", {'expire_days': 0})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIsNone(r.json()['envelope']['expires_at'])
        self.assertEqual(self.post(f"/api/esign/envelopes/{env['id']}/links", {'active': False}).status_code, 409, 'only a completed envelope')
        link = self.smtp.link(self.smtp.to('jane@client.test')[-1])
        self.sign('jane@client.test')
        # A year later the completed documents are still available from the email link.
        later = time.time() + 365 * 86400
        with patch('esign.now', return_value=later):
            fresh = self.smtp.link(self.smtp.to('jane@client.test')[-1])
            from fastapi.testclient import TestClient
            import portal
            visitor = TestClient(portal.app, base_url='http://127.0.0.1:8765')
            self.assertEqual(visitor.get('/sign/' + fresh, follow_redirects=False).status_code, 303)
        self.login(self.u2)
        self.assertEqual(self.post(f"/api/esign/envelopes/{env['id']}/links", {'active': False}).status_code, 404, 'another user')
        self.login(self.u1)
        r = self.post(f"/api/esign/envelopes/{env['id']}/links", {'active': False})
        self.assertEqual(r.status_code, 200, r.text)
        visitor = TestClient(portal.app, base_url='http://127.0.0.1:8765')
        self.assertEqual(visitor.get('/sign/' + self.smtp.link(self.smtp.to('jane@client.test')[-1]), follow_redirects=False).status_code, 410)
        del link


class Preparing(Base):
    def test_suggestions_place_initials_on_every_page_and_the_signature_block(self):
        contract_id, client_id = self.contract(self.u1)
        self.login(self.u1)
        env = self.post('/api/esign/envelopes', {'entity': SCRIBE, 'client_id': client_id}).json()
        library = self.api.get(f"/api/esign/envelopes/{env['id']}/library").json()['items']
        generated = [i for i in library if i['source'] == 'generated']
        self.assertEqual(sorted(i['agreement'] for i in generated), ['baa', 'services'])
        self.assertTrue(all(i['usable'] and i['latest'] for i in generated))
        r = self.post(f"/api/esign/envelopes/{env['id']}/library", {'source': 'generated', 'id': generated[0]['id']})
        self.assertEqual(r.status_code, 200, r.text)
        env = r.json()
        self.assertEqual(env['documents'][0]['source'], 'generated')
        self.assertEqual(self.post(f"/api/esign/envelopes/{env['id']}/library", {'source': 'generated', 'id': generated[0]['id']}).status_code, 409)
        s = self.api.get(f"/api/esign/envelopes/{env['id']}/suggestions?signers=2").json()['documents'][env['documents'][0]['id']]
        self.assertEqual(len(s['initials']), 2)
        self.assertEqual(len(s['initials'][0]), env['documents'][0]['pages'])
        self.assertEqual(s['blocks'][0]['fields'][0]['kind'], 'signature')

        # Someone else's client, another company's agreement, or another team: refused.
        other_contract, other_client = self.contract(self.u2)
        with db.engine.connect() as c:
            other_doc = c.execute(select(db.documents.c.id).where(db.documents.c.contract_id == other_contract)).scalar()
        self.assertEqual(self.post(f"/api/esign/envelopes/{env['id']}/library", {'source': 'generated', 'id': other_doc}).status_code, 404)
        self.login(self.u2)
        self.assertEqual(self.api.get(f"/api/esign/envelopes/{env['id']}/suggestions").status_code, 404)
        self.assertEqual(self.api.get(f"/api/esign/envelopes/{env['id']}/library").status_code, 404)

    def test_uploaded_form_labels_become_fields(self):
        from reportlab.pdfgen import canvas
        out = io.BytesIO()
        pdf = canvas.Canvas(out)
        pdf.setFont('Helvetica', 11)
        pdf.drawString(60, 700, 'Effective Date of this order is set out above and the Name of the client appears below.')
        pdf.drawString(60, 600, 'Client Signature: ______________________')
        pdf.drawString(60, 570, 'Printed Name: __________________________')
        pdf.drawString(60, 540, 'Date: ________________')
        pdf.save()
        s = esign_pdf.suggest(out.getvalue(), False, 1)
        self.assertEqual([f['kind'] for f in s['blocks'][0]['fields']], ['signature', 'name', 'date_signed'])
        self.assertGreater(s['blocks'][0]['fields'][0]['x'], 140, 'beside the label, on the line')
        self.assertEqual(len(s['initials'][0]), 1)


class SharedClients(Base):
    def test_admins_and_the_super_admin_reuse_their_teams_client_instead_of_creating_a_copy(self):
        import clients
        self.login(self.u1)
        r = self.post('/api/clients/new', {'name': 'Shared Family Clinic', 'entity': SCRIBE})
        self.assertEqual(r.status_code, 200, r.text)
        mine = r.json()['id']
        for user in (self.admin, self.super):          # u1's Admin and the Super Admin see u1's client
            self.login(user)
            self.assertEqual(self.post('/api/clients/new', {'name': 'shared family clinic ', 'entity': SCRIBE}).status_code, 409)
            with db.engine.begin() as c:
                self.assertEqual(clients.ensure_client(c, user, 'Shared Family Clinic', SCRIBE), mine, 'bulk send links the existing client')
        self.login(self.super)
        r = self.post('/api/clients', {'client_name': 'Shared Family Clinic', 'address': '1 Main St, Reno, NV 89501', 'entity': SCRIBE,
                                       'effective_date': '2026-10-24', 'document_types': ['baa'], 'output_root': str(self.root)})
        self.assertEqual((r.status_code, r.json().get('id')), (200, mine), r.text)
        with db.engine.connect() as c:
            count = c.execute(select(func.count()).select_from(db.clients).where(db.clients.c.name_key == clients.name_key('Shared Family Clinic'))).scalar()
        self.assertEqual(count, 1)
        self.login(self.ehs_user)                       # someone outside the team cannot see it, so gets their own client
        r = self.post('/api/clients/new', {'name': 'Shared Family Clinic', 'entity': EHS})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertNotEqual(r.json()['id'], mine)


class HomeReportsActivity(Base):
    def setUp(self):
        super().setUp()
        self.connect_mailbox()
        self.contracts = {u['id']: self.contract(u) for u in (self.u1, self.u2)}
        for user in (self.u1, self.u2):
            self.login(user)
            env = self.post('/api/esign/envelopes', {'contract_id': self.contracts[user['id']][0]}).json()
            people = [dict(r, name='Jane Example', key=r['id']) for r in env['recipients']]
            fields = [dict(f, recipient=f['recipient_id']) for f in env['fields']]
            subject = '=HYPERLINK("http://evil.test","x")' if user is self.u1 else 'Plain subject'
            self.post(f"/api/esign/envelopes/{env['id']}/draft", {'recipients': people, 'fields': fields, 'subject': subject})
            self.assertEqual(self.post(f"/api/esign/envelopes/{env['id']}/send").status_code, 200)
        self.drain()

    def test_home_counts_follow_the_scope(self):
        self.login(self.u1)
        mine = self.api.get('/api/home').json()
        self.assertEqual(mine['status']['waiting'], 1)
        self.assertEqual(mine['totals']['sent'], 1)
        self.assertEqual(mine['totals']['pending_signatures'], 1)
        self.login(self.admin)
        self.assertEqual(self.api.get('/api/home').json()['status']['waiting'], 2)
        self.login(self.ehs_user)
        self.assertEqual(self.api.get('/api/home').json()['status']['waiting'], 0)
        self.assertEqual(self.api.get('/api/home?entity=' + SCRIBE).status_code, 403)
        insights = self.api.get('/api/insights?tz=America/New_York').json()
        self.assertEqual(insights['totals']['sent'], 0)

    def test_reports_screen_excel_and_csv(self):
        self.login(self.u1)
        catalog = self.api.get('/api/reports/catalog').json()
        self.assertNotIn('activity', [d['key'] for d in catalog['datasets']], 'users do not get the activity report')
        self.assertEqual(self.post('/api/reports/run', {'dataset': 'activity'}).status_code, 403)
        self.assertEqual(self.post('/api/reports/run', {'dataset': 'users'}).status_code, 403)
        r = self.post('/api/reports/run', {'dataset': 'envelopes', 'filters': {'tz': 'America/Chicago'}})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['total'], 1, 'only my envelopes')
        r = self.post('/api/reports/export', {'dataset': 'envelopes', 'columns': ['ref', 'subject', 'status', 'sent'], 'format': 'xlsx'})
        self.assertEqual(r.status_code, 200)
        sheet = load_workbook(io.BytesIO(r.content))['Envelopes']
        self.assertEqual([c.value for c in sheet[1]], ['Envelope ID', 'Subject', 'Status', 'Sent'])
        self.assertEqual(sheet['B2'].data_type, 's', 'text from data is never a formula')
        self.assertTrue(str(sheet['B2'].value).startswith('=HYPERLINK'))
        csv_text = self.post('/api/reports/export', {'dataset': 'envelopes', 'columns': ['subject'], 'format': 'csv'}).content.decode('utf-8-sig')
        self.assertIn("'=HYPERLINK", csv_text)
        self.login(self.admin)
        self.assertEqual(self.post('/api/reports/run', {'dataset': 'envelopes'}).json()['total'], 2, "the admin's team")
        users = self.post('/api/reports/run', {'dataset': 'users', 'columns': ['email']}).json()
        emails = {row[0] for row in users['rows']}
        self.assertNotIn(self.ehs_user['email'], emails)
        self.assertIn(self.u1['email'], emails)
        # Saved reports belong to their owner.
        self.assertEqual(self.post('/api/reports/saved', {'name': 'Waiting', 'definition': {'dataset': 'envelopes', 'filters': {'status': 'waiting'}}}).status_code, 200)
        saved = self.api.get('/api/reports/catalog').json()['saved']
        self.login(self.u1)
        self.assertEqual(self.post(f"/api/reports/saved/{saved[0]['id']}/delete").status_code, 404)

    def test_activity_log_records_who_did_what_and_where_from(self):
        self.login(self.u1)
        client_id = self.contracts[self.u1['id']][1]
        self.api.get(f'/api/clients/{client_id}/profile')
        env = self.api.get('/api/esign/envelopes?box=waiting').json()['envelopes'][0]
        doc = self.api.get(f"/api/esign/envelopes/{env['id']}").json()['documents'][0]
        self.assertEqual(self.api.get(f"/api/esign/documents/{doc['id']}/file").status_code, 200)
        # The signer opens and signs.
        client, base, state = self.signer(self.smtp.link(self.smtp.to('jane@client.test')[0]))
        self.assertEqual(self.complete_signing(client, base, state).status_code, 200)
        self.assertEqual(self.api.get('/api/activity').status_code, 403, 'users cannot read the activity log')
        self.login(self.admin)
        log = self.api.get('/api/activity').json()['entries']
        actions = {(e['who'], e['action']) for e in log}
        self.assertIn((self.u1['name'], 'Opened a client profile'), actions)
        self.assertIn((self.u1['name'], 'Downloaded an envelope document'), actions)
        self.assertIn((self.u1['name'], 'Signed in'), actions)
        self.assertTrue(any(e['type'] == 'signer' and e['action'] == 'Signed' for e in log))
        signed_in = next(e for e in log if e['action'] == 'Signed in')
        self.assertEqual(signed_in['ip'], 'testclient')
        self.assertEqual(signed_in['browser'], 'Automated client')
        only_u2 = self.api.get('/api/activity?user=' + self.u2['id']).json()['entries']
        self.assertTrue(only_u2 and all(e['user_id'] == self.u2['id'] for e in only_u2))
        r = self.post('/api/activity/export', {'format': 'xlsx', 'filters': {}})
        self.assertEqual(r.status_code, 200)
        self.assertGreater(load_workbook(io.BytesIO(r.content))['Activity log'].max_row, 5)
        # The client's own activity has the portal and the signer's actions.
        timeline = self.api.get(f'/api/clients/{client_id}/activity').json()['activity']
        kinds = {e['type'] for e in timeline}
        self.assertEqual(kinds, {'user', 'signer'} | ({'system'} & kinds))
        # Another company's user sees none of it.
        self.login(self.ehs_user)
        self.assertEqual(self.api.get(f'/api/clients/{client_id}/activity').status_code, 404)


class MigrationData(unittest.TestCase):
    def test_the_old_default_expiry_is_removed_and_a_chosen_one_kept(self):
        import tempfile
        from alembic import command
        from sqlalchemy import create_engine, text
        import migrate
        with tempfile.TemporaryDirectory() as folder:
            engine = create_engine('sqlite:///' + folder + '/m.db')
            with engine.begin() as c:
                command.upgrade(migrate.config(c), '0003_docusign_parity')
                now = time.time()
                c.execute(text("INSERT INTO users (id,email,name,password,role,companies,active,reset_required,created) VALUES ('u','a@b.c','A','x','user','[]',1,0,:t)"), {'t': now})
                for ref, days in (('A', 30), ('B', 10)):
                    c.execute(text("INSERT INTO esign_envelopes (id,ref,owner_id,entity,subject,message,status,settings,version,created,updated,expires_at,expiry_warned,"
                                   "completion_attempts) VALUES (:id,:id,'u','SCRIBE ALIGN LLC','S','','sent',:s,1,:t,:t,:e,0,0)"),
                              {'id': ref, 's': '{"expire_days": %d}' % days, 't': now, 'e': now + days * 86400})
            with engine.begin() as c:
                command.upgrade(migrate.config(c), 'head')
            with engine.connect() as c:
                rows = dict(c.execute(text('SELECT id, expires_at FROM esign_envelopes')).all())
            self.assertIsNone(rows['A'])
            self.assertIsNotNone(rows['B'])
            engine.dispose()


if __name__ == '__main__':
    unittest.main()
